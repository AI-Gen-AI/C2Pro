"""Bounded readers for untrusted WBS source files: ``.xlsx``, ``.csv`` and ``wbs-import/v1`` JSON.

A reader never executes anything: no macros, no formula evaluation (a formula in a mapped column
is a BLOCKING error; the workbook's cached/calculated values are never read), no external links.
It only turns the immutable bytes into deterministic rows keyed by their source location.

Schedule columns (dates, durations, predecessors, ...) and cost columns are NEVER mapped into WBS
semantics: a schedule activity is not a WBS node and a budget line is not a WBS node. They are
reported as ignored (WARNING) in tabular files and rejected in the strict JSON contract.
"""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any

from src.wbs.domain.digest import DICTIONARY_LIST_KEYS
from src.wbs.imports.config import CSV_DELIMITERS
from src.wbs.imports.contracts import (
    JSON_IMPORT_CONTRACT,
    MAX_CELL_CHARS,
    MAX_COLUMNS,
    MAX_JSON_DEPTH,
    MAX_ROWS,
    MAX_SHEETS,
    MAX_XLSX_MEMBERS,
    MAX_XLSX_UNCOMPRESSED_BYTES,
    Diagnostic,
    Severity,
)

DICTIONARY_FIELDS = ("scope_statement", *DICTIONARY_LIST_KEYS)
STRUCTURAL_FIELDS = ("id", "parent_id", "code", "parent_code", "level", "name", "control_level", "decomposition_kind")
MAPPED_FIELDS = frozenset({*STRUCTURAL_FIELDS, *DICTIONARY_FIELDS})
HIERARCHY_FIELDS = frozenset({"parent_code", "level", "parent_id"})

# Normalized header -> canonical field. Anything else is reported and ignored.
_ALIASES: dict[str, str] = {
    "code": "code", "wbs_code": "code", "wbs": "code",
    "name": "name", "title": "name", "wbs_name": "name", "element": "name", "element_name": "name",
    "parent_code": "parent_code", "parent": "parent_code", "parent_wbs": "parent_code", "parent_wbs_code": "parent_code",
    "level": "level", "outline_level": "level", "outline": "level", "depth": "level",
    "id": "id", "external_id": "id", "source_id": "id", "key": "id",
    "parent_id": "parent_id", "parent_external_id": "parent_id", "parent_key": "parent_id",
    "control_level": "control_level", "decomposition_kind": "decomposition_kind",
    "description": "scope_statement",
    **{name: name for name in DICTIONARY_FIELDS},
}

SCHEDULE_COLUMNS = frozenset({
    "start", "finish", "end", "start_date", "end_date", "finish_date", "planned_start", "planned_end",
    "planned_finish", "actual_start", "actual_end", "actual_finish", "baseline_start", "baseline_end",
    "baseline_finish", "early_start", "early_finish", "late_start", "late_finish", "duration", "remaining_duration",
    "predecessors", "predecessor", "successors", "successor", "milestone", "float", "total_float", "free_float",
    "percent_complete", "pct_complete", "complete", "constraint", "constraint_date", "calendar", "activity_id",
    "task_id",
})
COST_COLUMNS = frozenset({
    "cost", "costs", "budget", "amount", "price", "unit_price", "unit_cost", "total_cost", "total", "quantity", "qty",
    "unit", "units", "currency", "rate", "value", "planned_value", "earned_value", "actual_cost", "bac", "eac",
})
# JSON keys that try to carry authority. A WBS source cannot approve, baseline or name canonical ids.
AUTHORITY_KEYS = frozenset({
    "approved", "approved_by", "approver", "approval", "baseline", "baseline_id", "baseline_no", "submitted",
    "submitted_by", "status", "node_id", "wbs_node_id", "canonical_id", "uuid", "authority", "change_set_id",
    "decided_by", "applied",
})


class FormulaCell(str):
    """An unevaluated spreadsheet formula (its text, never its calculated value)."""


@dataclass
class RawRow:
    source_ref: str
    source_ordinal: int
    values: dict[str, Any]
    formula_fields: frozenset[str] = frozenset()


@dataclass
class RawSource:
    sheet: str | None = None
    fields: frozenset[str] = frozenset()
    rows: list[RawRow] = field(default_factory=list)
    diagnostics: list[Diagnostic] = field(default_factory=list)


def _blocking(code: str, message: str, source_ref: str | None = None, field_name: str | None = None) -> Diagnostic:
    return Diagnostic(Severity.BLOCKING_ERROR, code, message, source_ref, field_name)


def _warning(code: str, message: str, source_ref: str | None = None, field_name: str | None = None) -> Diagnostic:
    return Diagnostic(Severity.WARNING, code, message, source_ref, field_name)


def normalize_header(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.replace("%", " percent ")
    return re.sub(r"[^0-9a-z]+", "_", text.strip().lower()).strip("_")


@dataclass
class _HeaderMap:
    columns: dict[int, str]
    diagnostics: list[Diagnostic]

    @property
    def fields(self) -> frozenset[str]:
        return frozenset(self.columns.values())

    @property
    def is_wbs(self) -> bool:
        return "name" in self.fields and bool(self.fields & HIERARCHY_FIELDS)


def map_header(cells: list[Any], where: str) -> _HeaderMap:
    columns: dict[int, str] = {}
    diagnostics: list[Diagnostic] = []
    seen: dict[str, int] = {}
    for index, cell in enumerate(cells):
        header = normalize_header(cell)
        if not header:
            continue
        canonical = _ALIASES.get(header)
        if canonical is None:
            if header in SCHEDULE_COLUMNS or header.endswith("_date"):
                diagnostics.append(_warning("SCHEDULE_FIELD_IGNORED", f"schedule column {header!r} is never WBS data",
                                            where, header))
            elif header in COST_COLUMNS:
                diagnostics.append(_warning("COST_FIELD_IGNORED", f"cost column {header!r} is never WBS authority",
                                            where, header))
            else:
                diagnostics.append(_warning("UNKNOWN_COLUMN_IGNORED", f"column {header!r} is not part of the WBS "
                                            "import contract", where, header))
            continue
        if canonical in seen:
            diagnostics.append(_blocking("DUPLICATE_COLUMN", f"two columns map to {canonical!r}", where, canonical))
            continue
        seen[canonical] = index
        columns[index] = canonical
    return _HeaderMap(columns, diagnostics)


def _require_wbs_header(header: _HeaderMap, where: str | None) -> list[Diagnostic]:
    problems: list[Diagnostic] = []
    if "name" not in header.fields:
        problems.append(_blocking("MISSING_NAME_COLUMN", "the WBS source has no name column", where, "name"))
    if not header.fields & HIERARCHY_FIELDS:
        problems.append(_blocking("MISSING_HIERARCHY_COLUMNS", "the WBS source has no parent_code, parent_id or level "
                                  "column: the hierarchy cannot be determined", where))
    return problems


# ---------------------------------------------------------------------------------------- XLSX
def _xlsx_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    text = str(value)
    return text if text.strip() else None


def read_xlsx(data: bytes, sheet_name: str | None) -> RawSource:
    source = RawSource()
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        members = archive.infolist()
    except (zipfile.BadZipFile, OSError, ValueError):
        source.diagnostics.append(_blocking("XLSX_UNREADABLE", "the file is not a readable .xlsx workbook"))
        return source
    if len(members) > MAX_XLSX_MEMBERS or sum(m.file_size for m in members) > MAX_XLSX_UNCOMPRESSED_BYTES:
        source.diagnostics.append(_blocking("XLSX_TOO_LARGE", "the workbook expands beyond the import bounds"))
        return source
    names = {m.filename.lower() for m in members}
    if any(name.endswith("vbaproject.bin") for name in names):
        source.diagnostics.append(_blocking("XLSX_MACROS_NOT_SUPPORTED", "macro-enabled workbooks are not imported"))
        return source
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except Exception:  # noqa: BLE001 - any parser failure of untrusted bytes is a BLOCKING diagnostic
        source.diagnostics.append(_blocking("XLSX_UNREADABLE", "the file is not a readable .xlsx workbook"))
        return source
    try:
        worksheets = list(workbook.worksheets)
        if len(worksheets) > MAX_SHEETS:
            source.diagnostics.append(_blocking("XLSX_TOO_MANY_SHEETS", f"more than {MAX_SHEETS} sheets"))
            return source
        headers: dict[str, _HeaderMap] = {}
        raw_headers: dict[str, list[Any]] = {}
        for ws in worksheets:
            first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            cells = list(first)
            while cells and _xlsx_text(cells[-1]) is None:
                cells.pop()
            raw_headers[ws.title] = cells
            headers[ws.title] = map_header(cells, f"sheet:{ws.title}/row:1")
        if sheet_name is not None:
            if sheet_name not in headers:
                source.diagnostics.append(_blocking("SHEET_NOT_FOUND", f"sheet {sheet_name!r} does not exist"))
                return source
            chosen = sheet_name
        else:
            candidates = [title for title, header in headers.items() if header.is_wbs]
            if not candidates:
                source.diagnostics.append(_blocking("NO_WBS_SHEET", "no sheet has a name column and a hierarchy column"))
                return source
            if len(candidates) > 1:
                source.diagnostics.append(_blocking(
                    "SHEET_SELECTION_REQUIRED",
                    f"several sheets match the WBS contract ({', '.join(sorted(candidates))}): choose one explicitly"))
                return source
            chosen = candidates[0]
        source.sheet = chosen
        where = f"sheet:{chosen}/row:1"
        if len(raw_headers[chosen]) > MAX_COLUMNS:
            source.diagnostics.append(_blocking("TOO_MANY_COLUMNS", f"more than {MAX_COLUMNS} columns", where))
            return source
        header = headers[chosen]
        source.diagnostics += header.diagnostics
        problems = _require_wbs_header(header, where)
        if problems:
            source.diagnostics += problems
            return source
        source.fields = header.fields
        ws = workbook[chosen]
        ordinal = 0
        for row_number, cells in enumerate(
                ws.iter_rows(min_row=2, max_col=max(len(raw_headers[chosen]), 1)), start=2):
            values: dict[str, Any] = {}
            formulas: set[str] = set()
            ref = f"sheet:{chosen}/row:{row_number}"
            for index, cell in enumerate(cells):
                field_name = header.columns.get(index)
                if field_name is None:
                    continue
                raw = getattr(cell, "value", None)
                if getattr(cell, "data_type", None) == "f":
                    formulas.add(field_name)
                    values[field_name] = FormulaCell(str(raw))
                    continue
                text = _xlsx_text(raw)
                if text is not None and len(text) > MAX_CELL_CHARS:
                    source.diagnostics.append(_blocking("CELL_TOO_LONG", f"more than {MAX_CELL_CHARS} characters",
                                                        ref, field_name))
                    text = text[:MAX_CELL_CHARS]
                values[field_name] = text
            if not formulas and all(v is None for v in values.values()):
                continue
            ordinal += 1
            if ordinal > MAX_ROWS:
                source.diagnostics.append(_blocking("TOO_MANY_ROWS", f"more than {MAX_ROWS} WBS rows"))
                source.rows = []
                return source
            source.rows.append(RawRow(ref, ordinal, values, frozenset(formulas)))
        return source
    finally:
        workbook.close()


# ---------------------------------------------------------------------------------------- CSV
def _decode(data: bytes) -> str | None:
    if data.startswith(b"\xef\xbb\xbf"):
        data = data[3:]
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return None


def _first_record(text: str, delimiter: str) -> list[str] | None:
    try:
        return next(csv.reader(io.StringIO(text), delimiter=delimiter, strict=True), None)
    except csv.Error:
        return None


def read_csv(data: bytes, delimiter: str | None) -> RawSource:
    source = RawSource()
    text = _decode(data)
    if text is None:
        source.diagnostics.append(_blocking("CSV_ENCODING", "a WBS CSV must be UTF-8 encoded"))
        return source
    if "\x00" in text:
        source.diagnostics.append(_blocking("CSV_INVALID", "the CSV contains NUL bytes"))
        return source
    if delimiter is None:
        matching = [d for d in CSV_DELIMITERS
                    if (first := _first_record(text, d)) is not None and map_header(list(first), "row:1").is_wbs]
        if not matching:
            source.diagnostics.append(_blocking("CSV_HEADER_NOT_RECOGNIZED", "no delimiter yields a header with a name "
                                                "and a hierarchy column", "row:1"))
            return source
        if len(matching) > 1:
            source.diagnostics.append(_blocking("CSV_DELIMITER_AMBIGUOUS", "the delimiter is ambiguous: set "
                                                "parse_config.delimiter", "row:1"))
            return source
        delimiter = matching[0]
    try:
        records = list(csv.reader(io.StringIO(text), delimiter=delimiter, strict=True))
    except csv.Error:
        source.diagnostics.append(_blocking("CSV_INVALID", "the CSV is malformed"))
        return source
    if not records:
        source.diagnostics.append(_blocking("EMPTY_IMPORT", "the CSV has no header"))
        return source
    header_cells = list(records[0])
    while header_cells and not header_cells[-1].strip():
        header_cells.pop()
    if len(header_cells) > MAX_COLUMNS:
        source.diagnostics.append(_blocking("TOO_MANY_COLUMNS", f"more than {MAX_COLUMNS} columns", "row:1"))
        return source
    header = map_header(header_cells, "row:1")
    source.diagnostics += header.diagnostics
    problems = _require_wbs_header(header, "row:1")
    if problems:
        source.diagnostics += problems
        return source
    source.fields = header.fields
    ordinal = 0
    for record_number, record in enumerate(records[1:], start=2):
        ref = f"row:{record_number}"
        if any(cell.strip() for cell in record[len(header_cells):]):
            source.diagnostics.append(_warning("EXTRA_CELLS_IGNORED", "cells beyond the header are ignored", ref))
        values: dict[str, Any] = {}
        for index, field_name in header.columns.items():
            cell = record[index] if index < len(record) else ""
            if len(cell) > MAX_CELL_CHARS:
                source.diagnostics.append(_blocking("CELL_TOO_LONG", f"more than {MAX_CELL_CHARS} characters",
                                                    ref, field_name))
                cell = cell[:MAX_CELL_CHARS]
            values[field_name] = cell if cell.strip() else None  # a leading '=' is DATA, never a formula
        if all(v is None for v in values.values()):
            continue
        ordinal += 1
        if ordinal > MAX_ROWS:
            source.diagnostics.append(_blocking("TOO_MANY_ROWS", f"more than {MAX_ROWS} WBS rows"))
            source.rows = []
            return source
        source.rows.append(RawRow(ref, ordinal, values))
    return source


# ---------------------------------------------------------------------------------------- JSON
_JSON_NODE_KEYS = frozenset({"id", "parent_id", "code", "parent_code", "level", "name", "control_level",
                             "decomposition_kind", "dictionary"})
_JSON_STRING_KEYS = ("id", "parent_id", "code", "parent_code", "name", "control_level", "decomposition_kind")


class _DuplicateKeyError(ValueError):
    pass


def _no_duplicate_keys(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for key, value in pairs:
        if key in out:
            raise _DuplicateKeyError(key)
        out[key] = value
    return out


def _reject_constant(name: str) -> Any:
    raise ValueError(f"non-finite number {name}")


def json_depth(text: str) -> int:
    """Maximum container nesting, computed without recursion (strings are skipped)."""
    depth = deepest = 0
    in_string = escaped = False
    for char in text:
        if in_string:
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == '"':
                in_string = False
        elif char == '"':
            in_string = True
        elif char in "[{":
            depth += 1
            deepest = max(deepest, depth)
        elif char in "]}":
            depth -= 1
    return deepest


def read_json(data: bytes) -> RawSource:
    source = RawSource()
    text = _decode(data)
    if text is None:
        source.diagnostics.append(_blocking("JSON_INVALID", "a WBS JSON must be UTF-8 encoded"))
        return source
    if json_depth(text) > MAX_JSON_DEPTH:
        source.diagnostics.append(_blocking("JSON_TOO_DEEP", f"the JSON nests deeper than {MAX_JSON_DEPTH} levels"))
        return source
    try:
        document = json.loads(text, object_pairs_hook=_no_duplicate_keys, parse_constant=_reject_constant)
    except _DuplicateKeyError as exc:
        source.diagnostics.append(_blocking("JSON_DUPLICATE_KEY", f"duplicate key {exc}"))
        return source
    except (ValueError, RecursionError):
        source.diagnostics.append(_blocking("JSON_INVALID", "the file is not valid JSON"))
        return source
    if not isinstance(document, dict):
        source.diagnostics.append(_blocking("JSON_SCHEMA", f"a {JSON_IMPORT_CONTRACT} document is an object"))
        return source
    unknown = sorted(set(document) - {"format", "nodes"})
    if unknown:
        source.diagnostics.append(_blocking("JSON_UNKNOWN_FIELD", f"unknown top-level keys {unknown}"))
    if document.get("format") != JSON_IMPORT_CONTRACT:
        source.diagnostics.append(_blocking("UNSUPPORTED_CONTRACT_VERSION",
                                            f"only {JSON_IMPORT_CONTRACT!r} is imported", None, "format"))
        return source
    nodes = document.get("nodes")
    if not isinstance(nodes, list):
        source.diagnostics.append(_blocking("JSON_SCHEMA", "'nodes' must be a list", None, "nodes"))
        return source
    if len(nodes) > MAX_ROWS:
        source.diagnostics.append(_blocking("TOO_MANY_ROWS", f"more than {MAX_ROWS} WBS nodes"))
        return source
    fields: set[str] = set()
    for index, node in enumerate(nodes):
        ref = f"/nodes/{index}"
        if not isinstance(node, dict):
            source.diagnostics.append(_blocking("JSON_SCHEMA", "a node is an object", ref))
            continue
        values: dict[str, Any] = {}
        for key, value in node.items():
            normalized = normalize_header(key)
            if key in AUTHORITY_KEYS:
                source.diagnostics.append(_blocking("AUTHORITY_FIELD_REJECTED", f"{key!r} cannot be imported: a WBS "
                                                    "source carries no approval, baseline or canonical identity",
                                                    ref, key))
            elif key not in _JSON_NODE_KEYS:
                code = ("SCHEDULE_FIELD_REJECTED" if normalized in SCHEDULE_COLUMNS or normalized.endswith("_date")
                        else "COST_FIELD_REJECTED" if normalized in COST_COLUMNS else "JSON_UNKNOWN_FIELD")
                source.diagnostics.append(_blocking(code, f"{key!r} is not part of {JSON_IMPORT_CONTRACT}", ref, key))
            elif key in _JSON_STRING_KEYS and value is not None and not isinstance(value, str):
                source.diagnostics.append(_blocking("FIELD_TYPE_INVALID", f"{key!r} must be a string", ref, key))
            elif key == "level" and value is not None and (isinstance(value, bool) or not isinstance(value, int)):
                source.diagnostics.append(_blocking("FIELD_TYPE_INVALID", "'level' must be an integer", ref, key))
            elif key == "dictionary" and value is not None and not isinstance(value, dict):
                source.diagnostics.append(_blocking("FIELD_TYPE_INVALID", "'dictionary' must be an object", ref, key))
            elif isinstance(value, str) and len(value) > MAX_CELL_CHARS:
                source.diagnostics.append(_blocking("CELL_TOO_LONG", f"more than {MAX_CELL_CHARS} characters", ref, key))
            else:
                values[key] = (value.strip() or None) if isinstance(value, str) else value
                fields.add(key)
        source.rows.append(RawRow(ref, index + 1, values))
    if "dictionary" in fields:
        fields.discard("dictionary")
        fields |= set(DICTIONARY_FIELDS)
    source.fields = frozenset(fields)
    if nodes and not any(d.severity is Severity.BLOCKING_ERROR for d in source.diagnostics):
        mapped = _HeaderMap(dict(enumerate(sorted(source.fields))), [])
        source.diagnostics += _require_wbs_header(mapped, None)
    return source
