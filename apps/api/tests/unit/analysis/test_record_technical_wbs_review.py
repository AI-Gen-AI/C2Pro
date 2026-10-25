"""ACC-01 S7: human review audit is bounded to trusted current source and approved WBS.

The DB integration + authenticated API route are separate acceptance gates; these
unit tests prove the service's ordered authority re-checks and recorded payload.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from uuid import UUID

import docx
import pytest

from src.analysis.application import record_technical_wbs_review as target
from src.analysis.domain.technical_wbs_link_candidate import TechnicalWBSLinkCandidate
from src.temporal.domain.current_revision import CurrentRevision, CurrentRevisionStatus
from src.temporal.domain.document_revision import DocumentRevision
from src.temporal.domain.project_event import ProjectEvent
from src.wbs.domain.governance import AuthorityState, WBSAuthority

TENANT = UUID("11111111-1111-1111-1111-111111111111")
PROJECT = UUID("22222222-2222-2222-2222-222222222222")
DOC = UUID("33333333-3333-3333-3333-333333333333")
REV = UUID("44444444-4444-4444-4444-444444444444")
CLAUSE = UUID("55555555-5555-5555-5555-555555555555")
NODE = UUID("66666666-6666-6666-6666-666666666666")
BASE = UUID("77777777-7777-7777-7777-777777777777")
ACTOR = UUID("88888888-8888-8888-8888-888888888888")
NOW = datetime(2026, 10, 25, tzinfo=UTC)
SOURCE = "Section 4.1.2 Concrete shall exceed 35 MPa."
EVENTS: list[ProjectEvent] = []


@dataclass
class _Source:
    async def read_verified_text(self, revision: DocumentRevision) -> str:
        assert revision.revision_id == REV
        return SOURCE


class _Result:
    def __init__(self, value):
        self._value = value

    def scalars(self):
        return self

    def first(self):
        return self._value


class _Clause:
    id = CLAUSE
    tenant_id = TENANT
    project_id = PROJECT
    document_id = DOC
    revision_id = REV
    clause_code = "4.1.2"
    clause_type = None
    title = SOURCE
    full_text = SOURCE
    text_start_offset = 0
    text_end_offset = len(SOURCE)
    extracted_entities = {
        "technical_requirement": {"source_section": "4.1.2", "verification_status": "unverified"},
        "evidence_location": {"revision_id": str(REV)},
    }
    extraction_confidence = None
    extraction_model = "deterministic-technical-section-v1"
    manually_verified = False
    verified_at = None


class _Session:
    def __init__(self, *, user_ok: bool = True, clause_ok: bool = True) -> None:
        self.user_ok = user_ok
        self.clause_ok = clause_ok
        self.events = []
        self.project_lock_count = 0

    async def execute(self, stmt, params=None):
        if params and "project" in params:
            self.project_lock_count += 1
            return _Result((PROJECT,))
        # Existing audit rows, if any, are supplied via self.events.
        return _Result(self.events[0] if self.events else None)

    async def scalar(self, query, params=None):
        if params and "project" in params:
            self.project_lock_count += 1
            return PROJECT
        if hasattr(query, "column_descriptions"):
            entity = query.column_descriptions[0].get("entity")
            if entity is target.User:
                return ACTOR if self.user_ok else None
            if entity is target.ClauseORM:
                return _Clause() if self.clause_ok else None
        raise AssertionError(f"unrecognized read: {query!s}")


def _revision():
    return DocumentRevision(
        revision_id=REV,
        document_id=DOC,
        project_id=PROJECT,
        tenant_id=TENANT,
        rev_no=1,
        blob_hash="a" * 64,
        blob_key="revisions/technical.docx",
        valid_from=NOW,
        created_at=NOW,
    )


def _authority(*, baseline_id=BASE):
    return WBSAuthority(
        state=AuthorityState.APPROVED_BASELINE,
        baseline_id=baseline_id,
        baseline_no=1,
        tree_digest="a" * 64,
        applied_at=NOW,
        open_change_sets=0,
        legacy_node_count=0,
    )


def _patch(monkeypatch, session, *, revision_status=CurrentRevisionStatus.TRUSTED, baseline=BASE):
    async def _gate(*args, **kwargs):
        assert kwargs["node_ids"] == {NODE}
        return _authority(baseline_id=baseline)

    class _RevisionReader:
        def __init__(self, db):
            assert db is session

        async def get_by_id(self, revision_id, tenant_id):
            assert (revision_id, tenant_id) == (REV, TENANT)
            return _revision()

    class _CurrentReader:
        def __init__(self, db):
            assert db is session

        async def resolve(self, *, tenant_id, document_id):
            assert (tenant_id, document_id) == (TENANT, DOC)
            return CurrentRevision(
                document_id=DOC,
                revision_id=REV,
                status=revision_status,
                single_revision=True,
            )

    class _EventWriter:
        def __init__(self, db):
            assert db is session

        async def append(self, event):
            assert event.event_type == target.REVIEW_EVENT
            EVENTS.append(event)
            session.events.append(
                type("EventRow", (), {
                    "event_id": event.event_id,
                    "actor": event.actor,
                    "payload": event.payload,
                })()
            )
            return event

    monkeypatch.setattr(target, "require_linkable_wbs_nodes", _gate)
    monkeypatch.setattr(target, "SqlAlchemyDocumentRevisionRepository", _RevisionReader)
    monkeypatch.setattr(target, "SqlAlchemyCurrentRevisionResolver", _CurrentReader)
    monkeypatch.setattr(target, "SqlAlchemyProjectEventRepository", _EventWriter)
    monkeypatch.setattr(
        target,
        "preview_technical_wbs_link",
        lambda **_kwargs: TechnicalWBSLinkCandidate(
            tenant_id=TENANT,
            project_id=PROJECT,
            technical_document_id=DOC,
            technical_revision_id=REV,
            source_blob_hash="a" * 64,
            source_clause_id=CLAUSE,
            source_section="4.1.2",
            source_start_offset=0,
            source_end_offset=len(SOURCE),
            proposed_wbs_baseline_id=BASE,
            proposed_wbs_node_id=NODE,
        ),
    )


async def _review(session, **overrides):
    args = {
        "tenant_id": TENANT,
        "project_id": PROJECT,
        "actor_user_id": ACTOR,
        "clause_id": CLAUSE,
        "expected_revision_id": REV,
        "expected_blob_hash": "a" * 64,
        "expected_wbs_baseline_id": BASE,
        "proposed_wbs_node_id": NODE,
        "decision": target.TechnicalMappingDecision.CONFIRMED_FOR_REVIEW,
        "idempotency_key": "review-key-A",
        "reason": "Source and target verified by project engineer",
    }
    args.update(overrides)
    return await target.RecordTechnicalWBSReviewService(
        session=session, source_reader=_Source()
    ).execute(**args)


@pytest.mark.asyncio
async def test_writes_only_append_audit_never_canonical_link(monkeypatch) -> None:
    EVENTS.clear()
    session = _Session()
    _patch(monkeypatch, session)
    result = await _review(session)
    assert result.replayed is False
    assert result.candidate.is_approved_link is False
    assert session.project_lock_count == 1
    assert len(EVENTS) == 1
    audit = EVENTS[0]
    assert audit.actor == f"user:{ACTOR}"
    assert audit.source_revision_id == REV
    assert audit.payload["authority"] == "AUDIT_ONLY_NOT_CANONICAL_LINK"
    assert audit.payload["technical_assessment"] == "NOT_ASSESSED"
    assert audit.payload["wbs_node_id"] == str(NODE)
    assert audit.payload["decision"] == "CONFIRMED_FOR_REVIEW"
    assert audit.confidence is None


@pytest.mark.asyncio
async def test_exact_replay_is_idempotent_and_conflict_is_refused(monkeypatch) -> None:
    EVENTS.clear()
    session = _Session()
    _patch(monkeypatch, session)
    first = await _review(session)
    again = await _review(session)
    assert again.replayed
    assert again.event_id == first.event_id
    assert len(EVENTS) == 1
    with pytest.raises(target.TechnicalMappingReviewError, match="already reviewed"):
        await _review(session, decision=target.TechnicalMappingDecision.REJECTED)
    assert len(EVENTS) == 1


@pytest.mark.asyncio
async def test_untrusted_source_fails_without_audit(monkeypatch) -> None:
    EVENTS.clear()
    session = _Session()
    _patch(monkeypatch, session, revision_status=CurrentRevisionStatus.SINGLE_REVISION)
    with pytest.raises(target.TechnicalMappingReviewError, match="no approved trusted revision"):
        await _review(session)
    assert EVENTS == []


@pytest.mark.asyncio
async def test_stale_baseline_fails_without_audit(monkeypatch) -> None:
    EVENTS.clear()
    session = _Session()
    _patch(monkeypatch, session, baseline=UUID("99999999-9999-9999-9999-999999999999"))
    with pytest.raises(target.TechnicalMappingReviewError, match="baseline changed"):
        await _review(session)
    assert EVENTS == []


@pytest.mark.asyncio
async def test_inactive_or_unverified_reviewer_fails_without_audit(monkeypatch) -> None:
    EVENTS.clear()
    session = _Session(user_ok=False)
    _patch(monkeypatch, session)
    with pytest.raises(target.TechnicalMappingReviewError, match="human reviewer"):
        await _review(session)
    assert EVENTS == []


@pytest.mark.asyncio
async def test_missing_clause_fails_without_audit(monkeypatch) -> None:
    EVENTS.clear()
    session = _Session(clause_ok=False)
    _patch(monkeypatch, session)
    with pytest.raises(target.TechnicalMappingReviewError, match="technical clause"):
        await _review(session)
    assert EVENTS == []


@pytest.mark.asyncio
async def test_invalid_action_or_idempotency_is_rejected_before_sql(monkeypatch) -> None:
    session = _Session()
    _patch(monkeypatch, session)
    with pytest.raises(target.TechnicalMappingReviewError, match="unknown human mapping decision"):
        await _review(session, decision="approved")
    with pytest.raises(target.TechnicalMappingReviewError, match="idempotency key"):
        await _review(session, idempotency_key="short")
    assert session.project_lock_count == 0


@pytest.mark.asyncio
async def test_verified_docx_reader_matches_worker_source_offsets(tmp_path: Path) -> None:
    """The production adapter parses SHA-pinned bytes in worker text-block order."""
    doc = docx.Document()
    doc.add_paragraph("Section 4.1.2 Concrete shall exceed 35 MPa.")
    doc.add_paragraph("Section 4.2 Steel must meet B500S.")
    path = tmp_path / "spec.docx"
    doc.save(path)
    content = path.read_bytes()
    digest = hashlib.sha256(content).hexdigest()

    class _Storage:
        async def download_object(self, key: str) -> Path:
            assert key == "revisions/technical.docx"
            return path

    rev = _revision().model_copy(update={"blob_hash": digest})
    reader = target.ImmutableDocxTechnicalTextReader(storage=_Storage())  # type: ignore[arg-type]
    parsed = await reader.read_verified_text(rev)
    assert parsed == (
        "Section 4.1.2 Concrete shall exceed 35 MPa.\n\n"
        "Section 4.2 Steel must meet B500S."
    )


@pytest.mark.asyncio
async def test_verified_docx_reader_denies_wrong_hash(tmp_path: Path) -> None:
    doc = docx.Document()
    doc.add_paragraph("Section 4.1.2 Concrete shall exceed 35 MPa.")
    path = tmp_path / "spec.docx"
    doc.save(path)

    class _Storage:
        async def download_object(self, key: str) -> Path:
            return path

    reader = target.ImmutableDocxTechnicalTextReader(storage=_Storage())  # type: ignore[arg-type]
    with pytest.raises(ValueError, match="immutable revision blob hash mismatch"):
        await reader.read_verified_text(_revision())


@pytest.mark.asyncio
async def test_unsupported_source_format_fails_closed() -> None:
    class _Storage:
        async def download_object(self, key: str):
            raise AssertionError("must not download unsupported format")

    reader = target.ImmutableDocxTechnicalTextReader(storage=_Storage())  # type: ignore[arg-type]
    rev = _revision().model_copy(update={"blob_key": "revisions/technical.pdf"})
    with pytest.raises(target.TechnicalMappingReviewError, match="format is not supported"):
        await reader.read_verified_text(rev)
