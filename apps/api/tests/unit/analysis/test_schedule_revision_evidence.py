"""PC-4: real XLSX, pinned revision, WBS identity and truthful date regressions.

These tests do NOT certify an approved schedule, a critical path or a live demo.
"""

from __future__ import annotations

from dataclasses import replace
from datetime import UTC, datetime
from hashlib import sha256
from pathlib import Path
import threading
from uuid import UUID

import openpyxl
import pytest

from src.analysis.application.extract_schedule_revision_evidence_use_case import (
    ExtractScheduleRevisionEvidenceUseCase,
)
from src.analysis.domain.schedule_activity_evidence import (
    InvalidScheduleEvidenceError,
    ScheduleDateState,
    observed_date,
)
from src.documents.adapters.parsers.excel_file_parser import (
    ExcelFileParser,
    ExcelParsingError,
)
from src.documents.domain.models import Document, DocumentStatus, DocumentType
from src.temporal.domain.document_revision import DocumentRevision

TENANT = UUID("11111111-1111-1111-1111-111111111111")
PROJECT = UUID("22222222-2222-2222-2222-222222222222")
DOCUMENT = UUID("33333333-3333-3333-3333-333333333333")
REVISION = UUID("44444444-4444-4444-4444-444444444444")
BASELINE = UUID("55555555-5555-5555-5555-555555555555")
NODE = UUID("66666666-6666-6666-6666-666666666666")
NOW = datetime(2026, 10, 17, tzinfo=UTC)


def _xlsx(path: Path, *, explicit_node: bool = True) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Gantt V1"
    sheet.append([
        "Activity ID", "Task", "Start Date", "End Date", "Duration",
        "WBS", "Predecessors", "WBS Node ID", "Milestone ID",
    ])
    sheet.append([
        "ACT-01", "Foundation", "2026-11-01", "2026-11-05",
        "4", "1.2", None, str(NODE) if explicit_node else None, None,
    ])
    sheet.append([
        "ACT-02", "Concrete", "2026-11-06", "2026-11-09",
        "3", "1.2", "ACT-01", None, None,
    ])
    sheet.append([
        None, "Foundation complete", "2026-11-10", "2026-11-10",
        "0", "1.2", "ACT-02", None, "M-1",
    ])
    workbook.save(path)
    workbook.close()


def _document() -> Document:
    return Document(
        id=DOCUMENT,
        project_id=PROJECT,
        tenant_id=TENANT,
        filename="schedule.xlsx",
        document_type=DocumentType.SCHEDULE,
        upload_status=DocumentStatus.PARSED,
        file_format=".xlsx",
    )


def _revision(path: Path) -> DocumentRevision:
    return DocumentRevision(
        revision_id=REVISION,
        document_id=DOCUMENT,
        project_id=PROJECT,
        tenant_id=TENANT,
        rev_no=1,
        blob_hash=sha256(path.read_bytes()).hexdigest(),
        blob_key="revisions/schedule.xlsx",
        valid_from=NOW,
        created_at=NOW,
    )


class _Approved:
    def __init__(self) -> None:
        self.calls = 0

    async def require(
        self, *, tenant_id: UUID, project_id: UUID, node_ids: set[UUID]
    ) -> UUID:
        self.calls += 1
        assert tenant_id == TENANT and project_id == PROJECT
        assert node_ids == {NODE}
        return BASELINE


class _Denied:
    async def require(
        self, *, tenant_id: UUID, project_id: UUID, node_ids: set[UUID]
    ) -> UUID:
        raise InvalidScheduleEvidenceError("foreign or retired WBS node")



@pytest.mark.parametrize("include_evidence", [False, True])
@pytest.mark.asyncio
async def test_schedule_openpyxl_lifecycle_runs_off_event_loop(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, include_evidence: bool,
) -> None:
    """Both XLSX opens in evidence mode execute in a worker, not on the event loop."""
    path = tmp_path / "threaded-schedule.xlsx"
    _xlsx(path)
    original_open = openpyxl.load_workbook
    caller_thread = threading.current_thread()
    called_on: list[threading.Thread] = []

    def _record_open(*args: object, **kwargs: object):
        called_on.append(threading.current_thread())
        return original_open(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", _record_open)
    rows = await ExcelFileParser().parse_schedule(
        path, include_evidence=include_evidence
    )
    assert len(rows) == 3
    assert len(called_on) == (2 if include_evidence else 1)
    assert all(worker is not caller_thread for worker in called_on)


@pytest.mark.asyncio
async def test_schedule_legacy_parser_unchanged_and_source_evidence_opt_in(tmp_path: Path) -> None:
    path = tmp_path / "schedule.xlsx"
    _xlsx(path)
    parser = ExcelFileParser()
    legacy = await parser.parse_schedule(path)
    evidence = await parser.parse_schedule(path, include_evidence=True)

    assert len(legacy) == len(evidence) == 3
    assert set(legacy[0]) == {
        "task", "start_date", "end_date", "duration", "wbs", "predecessors",
    }
    assert evidence[0]["source_locator"] == "Gantt V1!2"
    assert evidence[0]["source_activity_id"] == "ACT-01"
    assert evidence[0]["wbs_node_id"] == str(NODE)
    assert evidence[2]["source_milestone_id"] == "M-1"
    # Legacy accidentally resolves the WBS Node ID column as generic "wbs"
    # when evidence aliases are off. Evidence mode must retain the explicit
    # source WBS *code* separately from the approved node UUID.
    assert legacy[0]["wbs"] == str(NODE)
    assert evidence[0]["wbs"] == "1.2"
    assert {key: evidence[0][key] for key in legacy[0] if key != "wbs"} == {
        key: value for key, value in legacy[0].items() if key != "wbs"
    }


@pytest.mark.asyncio
async def test_activities_and_milestone_are_read_only_revision_observations(tmp_path: Path) -> None:
    path = tmp_path / "schedule.xlsx"
    _xlsx(path)
    approved = _Approved()
    result = await ExtractScheduleRevisionEvidenceUseCase(
        parser=ExcelFileParser(), approved_scope_service=approved,
    ).execute(document=_document(), revision=_revision(path), source_path=path)
    assert len(result.rows) == 3
    first, second, milestone = result.rows
    assert approved.calls == 1
    assert first.identity == (REVISION, "Gantt V1!2")
    assert first.source_activity_id == "ACT-01"
    assert first.wbs_node_id == NODE
    assert first.approved_wbs_baseline_id == BASELINE
    assert first.source_wbs_code == "1.2"
    assert first.start_date.isoformat() == "2026-11-01"
    assert first.end_date.isoformat() == "2026-11-05"
    assert first.date_state is ScheduleDateState.SOURCE_DATES_OBSERVED
    assert second.wbs_node_id is None  # WBS code 1.2 alone is NOT node identity
    assert second.predecessors_source == "ACT-01"  # not a validated dependency edge
    assert milestone.source_milestone_id == "M-1"
    assert milestone.start_date == milestone.end_date
    assert result.time_assessment == "NOT_ASSESSED"


@pytest.mark.asyncio
async def test_schedule_source_code_alone_never_links_canonical_wbs(tmp_path: Path) -> None:
    path = tmp_path / "schedule.xlsx"
    _xlsx(path, explicit_node=False)
    approved = _Approved()
    result = await ExtractScheduleRevisionEvidenceUseCase(
        parser=ExcelFileParser(), approved_scope_service=approved,
    ).execute(document=_document(), revision=_revision(path), source_path=path)
    assert approved.calls == 0
    assert all(row.wbs_node_id is None for row in result.rows)
    assert all(row.approved_wbs_baseline_id is None for row in result.rows)
    assert result.rows[0].source_wbs_code == "1.2"


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("document_type", DocumentType.CONTRACT),
        ("id", UUID("77777777-7777-7777-7777-777777777777")),
        ("tenant_id", UUID("88888888-8888-8888-8888-888888888888")),
        ("project_id", UUID("99999999-9999-9999-9999-999999999999")),
    ],
)
@pytest.mark.asyncio
async def test_wrong_document_identity_rejected(
    tmp_path: Path, field: str, value: object,
) -> None:
    path = tmp_path / "schedule.xlsx"
    _xlsx(path)
    with pytest.raises(InvalidScheduleEvidenceError):
        await ExtractScheduleRevisionEvidenceUseCase(
            parser=ExcelFileParser(), approved_scope_service=_Approved(),
        ).execute(
            document=replace(_document(), **{field: value}),
            revision=_revision(path),
            source_path=path,
        )


@pytest.mark.asyncio
async def test_historical_xlsx_revision_survives_current_document_bc3_format(
    tmp_path: Path,
) -> None:
    """Mutable current document metadata cannot invalidate a pinned XLSX revision."""
    path = tmp_path / "schedule.xlsx"
    _xlsx(path)
    current_document = replace(_document(), file_format=".bc3", filename="schedule.bc3")
    result = await ExtractScheduleRevisionEvidenceUseCase(
        parser=ExcelFileParser(), approved_scope_service=_Approved(),
    ).execute(document=current_document, revision=_revision(path), source_path=path)
    assert len(result.rows) == 3
    assert result.rows[0].revision_id == REVISION


@pytest.mark.asyncio
async def test_revision_with_non_xlsx_object_key_is_rejected_even_if_document_says_xlsx(
    tmp_path: Path,
) -> None:
    path = tmp_path / "schedule.xlsx"
    _xlsx(path)
    with pytest.raises(InvalidScheduleEvidenceError, match="selected source revision"):
        await ExtractScheduleRevisionEvidenceUseCase(
            parser=ExcelFileParser(), approved_scope_service=_Approved(),
        ).execute(
            document=_document(),
            revision=_revision(path).model_copy(update={"blob_key": "revisions/schedule.bc3"}),
            source_path=path,
        )


@pytest.mark.asyncio
async def test_legacy_schedule_header_activity_id_still_maps_as_activity(tmp_path: Path) -> None:
    """Evidence-only ID aliases must not steal legacy 'activity' task matching."""
    path = tmp_path / "legacy-activity.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Activity ID", "Start Date", "End Date"])
    sheet.append(["A-101", "2026-11-01", "2026-11-05"])
    workbook.save(path)
    workbook.close()

    legacy = await ExcelFileParser().parse_schedule(path)
    assert legacy[0]["task"] == "A-101"


@pytest.mark.asyncio
async def test_modified_revision_blob_fails_before_parsing(tmp_path: Path) -> None:
    path = tmp_path / "schedule.xlsx"
    _xlsx(path)
    revision = _revision(path)
    path.write_bytes(path.read_bytes() + b"wrong-immutable-revision")
    with pytest.raises(ValueError, match="blob hash mismatch"):
        await ExtractScheduleRevisionEvidenceUseCase(
            parser=ExcelFileParser(), approved_scope_service=_Approved(),
        ).execute(document=_document(), revision=revision, source_path=path)


@pytest.mark.asyncio
async def test_unapproved_wbs_node_rejected_without_wbs_mutation(tmp_path: Path) -> None:
    path = tmp_path / "schedule.xlsx"
    _xlsx(path)
    with pytest.raises(InvalidScheduleEvidenceError, match="retired WBS"):
        await ExtractScheduleRevisionEvidenceUseCase(
            parser=ExcelFileParser(), approved_scope_service=_Denied(),
        ).execute(document=_document(), revision=_revision(path), source_path=path)


@pytest.mark.asyncio
async def test_replayed_activity_skips_changed_wbs_authority(tmp_path: Path) -> None:
    path = tmp_path / "schedule.xlsx"
    _xlsx(path)
    result = await ExtractScheduleRevisionEvidenceUseCase(
        parser=ExcelFileParser(), approved_scope_service=_Denied(),
    ).execute(
        document=_document(),
        revision=_revision(path),
        source_path=path,
        existing_row_identities={(REVISION, "Gantt V1!2")},
    )
    assert len(result.rows) == 2
    assert all(row.wbs_node_id is None for row in result.rows)


@pytest.mark.asyncio
async def test_duplicate_declared_activity_id_rejected(tmp_path: Path) -> None:
    path = tmp_path / "schedule.xlsx"
    _xlsx(path)
    workbook = openpyxl.load_workbook(path)
    workbook.active.cell(row=3, column=1).value = "ACT-01"
    workbook.save(path)
    workbook.close()
    with pytest.raises(InvalidScheduleEvidenceError, match="duplicate source activity id"):
        await ExtractScheduleRevisionEvidenceUseCase(
            parser=ExcelFileParser(), approved_scope_service=_Approved(),
        ).execute(document=_document(), revision=_revision(path), source_path=path)


@pytest.mark.asyncio
async def test_formula_date_rejected_even_when_source_hash_is_valid(tmp_path: Path) -> None:
    path = tmp_path / "schedule.xlsx"
    _xlsx(path)
    workbook = openpyxl.load_workbook(path)
    workbook.active.cell(row=2, column=3).value = "=TODAY()"
    workbook.save(path)
    workbook.close()
    with pytest.raises(ExcelParsingError, match="Unverified schedule date formula"):
        await ExtractScheduleRevisionEvidenceUseCase(
            parser=ExcelFileParser(), approved_scope_service=_Approved(),
        ).execute(document=_document(), revision=_revision(path), source_path=path)



@pytest.mark.parametrize(
    "formula_column", [1, 2, 3, 4, 5, 6, 7, 8, 9]
)
@pytest.mark.asyncio
async def test_schedule_evidence_rejects_formulas_in_every_consumed_field(
    tmp_path: Path, formula_column: int,
) -> None:
    """Verified XLSX bytes cannot attest cached task/ID/WBS/dependency values."""
    path = tmp_path / "schedule-all-fields.xlsx"
    _xlsx(path)
    workbook = openpyxl.load_workbook(path)
    workbook.active.cell(row=2, column=formula_column).value = '="unverified"'
    workbook.save(path)
    workbook.close()

    authority = _Approved()
    with pytest.raises(ExcelParsingError, match="Unverified schedule"):
        await ExtractScheduleRevisionEvidenceUseCase(
            parser=ExcelFileParser(), approved_scope_service=authority,
        ).execute(document=_document(), revision=_revision(path), source_path=path)
    assert authority.calls == 0


@pytest.mark.asyncio
async def test_schedule_evidence_does_not_skip_formula_only_source_rows(
    tmp_path: Path,
) -> None:
    """Formula without a cached result must fail closed, not vanish as a blank row."""
    path = tmp_path / "formula-blank-cache.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Task", "Start Date", "End Date"])
    sheet.append(['="hidden task"', None, None])
    workbook.save(path)
    workbook.close()

    with pytest.raises(ExcelParsingError, match="Unverified schedule source formula"):
        await ExtractScheduleRevisionEvidenceUseCase(
            parser=ExcelFileParser(), approved_scope_service=_Approved(),
        ).execute(document=_document(), revision=_revision(path), source_path=path)


def test_ambiguous_dates_and_timezone_are_never_inferred() -> None:
    from datetime import timedelta, timezone

    assert observed_date("01/02/2026") is None
    assert observed_date(46000) is None
    assert observed_date("2026-10-17T08:00:00") is None
    assert observed_date(datetime(2026, 10, 17, 8, 0)) is None
    assert observed_date(datetime(2026, 10, 17, tzinfo=timezone(timedelta(hours=2)))) is None


@pytest.mark.asyncio
async def test_ambiguous_dates_and_reversed_range_are_unassessed(tmp_path: Path) -> None:
    path = tmp_path / "schedule.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Task", "Start Date", "End Date"])
    sheet.append(["Ambiguous", "01/02/2026", "2026-10-03"])
    sheet.append(["Inverted", "2026-12-31", "2026-01-01"])
    workbook.save(path)
    workbook.close()

    result = await ExtractScheduleRevisionEvidenceUseCase(
        parser=ExcelFileParser(), approved_scope_service=_Approved(),
    ).execute(document=_document(), revision=_revision(path), source_path=path)
    assert [row.date_state for row in result.rows] == [
        ScheduleDateState.NOT_ASSESSED, ScheduleDateState.NOT_ASSESSED,
    ]
    assert [row.date_reason for row in result.rows] == [
        "missing_or_ambiguous_source_date", "source_end_before_start",
    ]
    assert result.time_assessment == "NOT_ASSESSED"
