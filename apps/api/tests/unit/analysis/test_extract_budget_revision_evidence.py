"""PC-3 follow-up: real spreadsheet rows become revision-bound evidence."""

from __future__ import annotations

import threading
from dataclasses import replace
from datetime import UTC, datetime
from decimal import Decimal
from hashlib import sha256
from pathlib import Path
from uuid import UUID

import openpyxl
import pytest

from src.analysis.application.extract_budget_revision_evidence_use_case import (
    ExtractBudgetRevisionEvidenceUseCase,
)
from src.analysis.domain.budget_row_evidence import (
    InvalidBudgetRowEvidenceError,
    UnapprovedWBSMappingError,
)
from src.documents.adapters.parsers.excel_file_parser import ExcelFileParser, ExcelParsingError
from src.documents.domain.models import Document, DocumentStatus, DocumentType
from src.temporal.domain.document_revision import DocumentRevision

TENANT = UUID("11111111-1111-1111-1111-111111111111")
PROJECT = UUID("22222222-2222-2222-2222-222222222222")
DOCUMENT = UUID("33333333-3333-3333-3333-333333333333")
REVISION = UUID("44444444-4444-4444-4444-444444444444")
BASELINE = UUID("55555555-5555-5555-5555-555555555555")
WBS_NODE = UUID("66666666-6666-6666-6666-666666666666")
NOW = datetime(2026, 10, 17, tzinfo=UTC)


def _budget_xlsx(path: Path) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Budget 2026"
    worksheet.append(
        ["Item", "Quantity", "Unit", "Unit Price", "Total", "Currency", "Tax Basis", "WBS Node ID"]
    )
    worksheet.append(["Concrete", 12, "m3", 100, 1200, "EUR", "NET", str(WBS_NODE)])
    workbook.save(path)


def _document() -> Document:
    return Document(
        id=DOCUMENT,
        tenant_id=TENANT,
        project_id=PROJECT,
        filename="budget.xlsx",
        document_type=DocumentType.BUDGET,
        upload_status=DocumentStatus.PARSED,
        file_format=".xlsx",
    )


def _revision(source: Path | None = None) -> DocumentRevision:
    return DocumentRevision(
        revision_id=REVISION,
        document_id=DOCUMENT,
        project_id=PROJECT,
        tenant_id=TENANT,
        rev_no=1,
        blob_hash=sha256(source.read_bytes()).hexdigest() if source else "a" * 64,
        blob_key="revisions/budget.xlsx",
        valid_from=NOW,
        created_at=NOW,
    )


class _ApprovedScopeService:
    async def require(self, *, tenant_id: UUID, project_id: UUID, node_ids: set[UUID]):
        assert tenant_id == TENANT
        assert project_id == PROJECT
        assert node_ids == {WBS_NODE}
        return BASELINE




@pytest.mark.asyncio
async def test_budget_evidence_rejects_missing_formula_workbook(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A source workbook without the formula view cannot certify evidence."""
    source = tmp_path / "missing-formula-view.xlsx"
    _budget_xlsx(source)
    original_open = openpyxl.load_workbook
    opened_formula_view = False

    def _missing_formula_view(*args: object, **kwargs: object):
        nonlocal opened_formula_view
        if kwargs.get("data_only") is False:
            opened_formula_view = True
            return None
        return original_open(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", _missing_formula_view)
    with pytest.raises(
        ExcelParsingError,
        match="Formula workbook unavailable for verified budget evidence",
    ):
        await ExcelFileParser().parse_budget(source, include_evidence=True)
    assert opened_formula_view


@pytest.mark.asyncio
async def test_budget_revision_hash_validation_runs_off_event_loop(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Hashing still verifies immutable bytes, but cannot block other coroutines."""
    from src.analysis.application import extract_budget_revision_evidence_use_case as module

    source = tmp_path / "hash-worker-budget.xlsx"
    _budget_xlsx(source)
    original_verify = module.verify_revision_file_hash
    hashing_threads: list[threading.Thread] = []
    caller_thread = threading.current_thread()

    def _record_hash(source_path: Path, blob_hash: str) -> None:
        hashing_threads.append(threading.current_thread())
        original_verify(source_path, blob_hash)

    monkeypatch.setattr(module, "verify_revision_file_hash", _record_hash)
    result = await ExtractBudgetRevisionEvidenceUseCase(
        parser=ExcelFileParser(), approved_scope_service=_ApprovedScopeService(),
    ).execute(
        document=_document(), revision=_revision(source), source_path=source,
        existing_row_identities=frozenset(),
    )
    assert len(result.rows) == 1
    assert len(hashing_threads) == 1
    assert hashing_threads[0] is not caller_thread


@pytest.mark.parametrize("include_evidence", [False, True])
@pytest.mark.asyncio
async def test_budget_workbooks_are_parsed_off_event_loop(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, include_evidence: bool,
) -> None:
    path = tmp_path / "threaded-budget.xlsx"
    _budget_xlsx(path)
    original_open = openpyxl.load_workbook
    caller = threading.current_thread()
    calls: list[threading.Thread] = []

    def _record_open(*args: object, **kwargs: object):
        calls.append(threading.current_thread())
        return original_open(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", _record_open)
    result = await ExcelFileParser().parse_budget(path, include_evidence=include_evidence)
    assert len(result) == 1
    assert len(calls) == (2 if include_evidence else 1)
    assert all(worker is not caller for worker in calls)


@pytest.mark.asyncio
async def test_budget_parser_legacy_shape_is_stable_with_opt_in_source_evidence(
    tmp_path: Path,
) -> None:
    """Existing analysis/BOM readers retain their original budget row fields."""
    source = tmp_path / "budget.xlsx"
    _budget_xlsx(source)
    parser = ExcelFileParser()

    [legacy] = await parser.parse_budget(source)
    [evidenced] = await parser.parse_budget(source, include_evidence=True)

    assert set(legacy) == {
        "item", "quantity", "unit_price", "total", "unit", "category"
    }
    assert evidenced["source_locator"] == "Budget 2026!2"
    assert evidenced["numeric_provenance"] == {
        "quantity": "12", "unit price": "100", "total": "1200"
    }
    assert evidenced["currency"] == "EUR"
    assert evidenced["tax_basis"] == "NET"
    assert evidenced["wbs_node_id"] == str(WBS_NODE)
    assert {key: evidenced[key] for key in legacy} == legacy

@pytest.mark.asyncio
async def test_extracts_a_real_xlsx_row_with_locator_provenance_and_approved_wbs(
    tmp_path: Path,
) -> None:
    """Removing immutable source, numeric, or WBS checks makes this row unsafe."""
    source = tmp_path / "budget.xlsx"
    _budget_xlsx(source)

    result = await ExtractBudgetRevisionEvidenceUseCase(
        parser=ExcelFileParser(), approved_scope_service=_ApprovedScopeService()
    ).execute(
        document=_document(),
        revision=_revision(source),
        source_path=source,
        existing_row_identities=frozenset(),
    )

    [row] = result.rows
    assert row.source_locator == "Budget 2026!2"
    assert row.quantity == 12
    assert row.unit_price == 100
    assert row.amount == 1200
    assert row.numeric_provenance == (("quantity", "12"), ("unit price", "100"), ("total", "1200"))
    assert row.currency == "EUR"
    assert row.tax_basis == "NET"
    assert row.wbs_baseline_id == BASELINE
    assert row.wbs_node_id == WBS_NODE
    assert result.assessment_reason == "approved_budget_comparison_missing"



@pytest.mark.asyncio
async def test_immutable_budget_evidence_preserves_large_decimal_text(
    tmp_path: Path,
) -> None:
    """Converting financial text via legacy float silently changes the value."""
    source = tmp_path / "precise.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Item", "Quantity", "Unit", "Unit Price", "Total", "Currency", "Tax Basis"])
    sheet.append([
        "Exact amount", "1", "each", "99999999999999.99", "99999999999999.99",
        "EUR", "NET"
    ])
    workbook.save(source)

    [legacy] = await ExcelFileParser().parse_budget(source)
    assert isinstance(legacy["total"], float)

    result = await ExtractBudgetRevisionEvidenceUseCase(
        parser=ExcelFileParser(), approved_scope_service=_ApprovedScopeService()
    ).execute(
        document=_document(), revision=_revision(source), source_path=source,
        existing_row_identities=frozenset(),
    )
    [row] = result.rows
    assert row.quantity == Decimal("1")
    assert row.unit_price == Decimal("99999999999999.99")
    assert row.amount == Decimal("99999999999999.99")
    assert row.numeric_provenance == (
        ("quantity", "'1'"),
        ("unit price", "'99999999999999.99'"),
        ("total", "'99999999999999.99'"),
    )


@pytest.mark.asyncio
async def test_immutable_budget_evidence_handles_es_decimal_source_text(
    tmp_path: Path,
) -> None:
    """Spanish decimal separators must not be rounded through binary floats."""
    source = tmp_path / "es-budget.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Item", "Quantity", "Unit", "Unit Price", "Total", "Currency", "Tax Basis"])
    sheet.append(["Concrete", "1.234,50", "m3", "12,30 EUR", "15.184,35 EUR", "EUR", "NET"])
    workbook.save(source)

    result = await ExtractBudgetRevisionEvidenceUseCase(
        parser=ExcelFileParser(), approved_scope_service=_ApprovedScopeService()
    ).execute(
        document=_document(), revision=_revision(source), source_path=source,
        existing_row_identities=frozenset(),
    )
    [row] = result.rows
    assert (row.quantity, row.unit_price, row.amount) == (
        Decimal("1234.50"), Decimal("12.30"), Decimal("15184.35")
    )


@pytest.mark.parametrize("formula_column", ["unit price", "total"])
@pytest.mark.asyncio
async def test_rejects_unverified_formula_cells_even_if_workbook_is_hash_verified(
    tmp_path: Path, formula_column: str
) -> None:
    """Hash pins bytes, but does not prove a formula was evaluated or its cache is fresh."""
    source = tmp_path / "formula.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Item", "Quantity", "Unit", "Unit Price", "Total", "Currency", "Tax Basis"])
    sheet.append(["Concrete", 12, "m3", 100, 1200, "EUR", "NET"])
    sheet.cell(row=2, column=4 if formula_column == "unit price" else 5).value = "=B2*4"
    workbook.save(source)

    with pytest.raises(ExcelParsingError, match="Unverified formula-derived budget evidence"):
        await ExtractBudgetRevisionEvidenceUseCase(
            parser=ExcelFileParser(), approved_scope_service=_ApprovedScopeService()
        ).execute(
            document=_document(), revision=_revision(source), source_path=source,
            existing_row_identities=frozenset(),
        )


@pytest.mark.parametrize(
    "formula_column", [1, 3, 6, 7, 8]
)
@pytest.mark.asyncio
async def test_rejects_cached_formulas_in_any_evidence_bearing_budget_field(
    tmp_path: Path, formula_column: int,
) -> None:
    """Hash-pinned but stale formula cells cannot attest WBS, amount or labels."""
    path = tmp_path / "formula-all-fields.xlsx"
    _budget_xlsx(path)
    workbook = openpyxl.load_workbook(path)
    workbook.active.cell(row=2, column=formula_column).value = '="stale"'
    workbook.save(path)
    workbook.close()
    with pytest.raises(ExcelParsingError, match="Unverified formula-derived budget evidence"):
        await ExtractBudgetRevisionEvidenceUseCase(
            parser=ExcelFileParser(), approved_scope_service=_ApprovedScopeService()
        ).execute(
            document=_document(), revision=_revision(path), source_path=path,
            existing_row_identities=frozenset(),
        )


@pytest.mark.parametrize("amount", ["0,125", "1,234", "25,1234"])
@pytest.mark.asyncio
async def test_rejects_ambiguous_single_comma_decimal_source_text(
    tmp_path: Path, amount: str,
) -> None:
    """Unknown locale must not silently magnify financial values by 1000x."""
    path = tmp_path / "budget-ambiguous.xlsx"
    _budget_xlsx(path)
    workbook = openpyxl.load_workbook(path)
    workbook.active.cell(row=2, column=2).value = amount
    workbook.save(path)
    workbook.close()
    with pytest.raises(ExcelParsingError, match="Ambiguous single-comma decimal"):
        await ExtractBudgetRevisionEvidenceUseCase(
            parser=ExcelFileParser(), approved_scope_service=_ApprovedScopeService()
        ).execute(
            document=_document(), revision=_revision(path), source_path=path,
            existing_row_identities=frozenset(),
        )


@pytest.mark.parametrize(
    "source",
    ["1.234", "0.125", "1e3", "12abc34", "1 234", "2+3", "1,234.56.7"],
)
def test_verified_budget_amount_rejects_ambiguous_or_unsupported_source_syntax(
    source: str,
) -> None:
    with pytest.raises(ExcelParsingError):
        ExcelFileParser._coerce_budget_decimal(source)


@pytest.mark.parametrize(
    ("source", "expected"),
    [
        ("1.234,50", Decimal("1234.50")),
        ("1,234.50", Decimal("1234.50")),
        ("12,30 EUR", Decimal("12.30")),
        ("99999999999999.99", Decimal("99999999999999.99")),
        ("(100.00)", Decimal("-100.00")),
        ("1.234.567", Decimal("1234567")),
        ("1,234,567", Decimal("1234567")),
    ],
)
def test_verified_budget_amount_preserves_exact_numeric_semantics(
    source: str, expected: Decimal,
) -> None:
    assert ExcelFileParser._coerce_budget_decimal(source) == expected


@pytest.mark.parametrize(
    ("unit_source", "total_source", "currency"),
    [
        ("$100", "$1200", "EUR"),
        ("EUR 100 USD", "1200", "EUR"),
        ("100", "£1200", "EUR"),
        ("€100", "€1200", None),
    ],
)
@pytest.mark.asyncio
async def test_budget_evidence_rejects_foreign_or_conflicting_currency_markers(
    tmp_path: Path, unit_source: str, total_source: str, currency: str | None,
) -> None:
    path = tmp_path / "conflicting-currency.xlsx"
    _budget_xlsx(path)
    workbook = openpyxl.load_workbook(path)
    sheet = workbook.active
    sheet.cell(row=2, column=4).value = unit_source
    sheet.cell(row=2, column=5).value = total_source
    sheet.cell(row=2, column=6).value = currency
    workbook.save(path)
    workbook.close()
    with pytest.raises(ExcelParsingError, match="currency marker conflicts"):
        await ExtractBudgetRevisionEvidenceUseCase(
            parser=ExcelFileParser(), approved_scope_service=_ApprovedScopeService()
        ).execute(
            document=_document(), revision=_revision(path), source_path=path,
            existing_row_identities=frozenset(),
        )


@pytest.mark.asyncio
async def test_budget_evidence_accepts_matching_explicit_currency_markers(
    tmp_path: Path,
) -> None:
    path = tmp_path / "matching-currency.xlsx"
    _budget_xlsx(path)
    workbook = openpyxl.load_workbook(path)
    sheet = workbook.active
    sheet.cell(row=2, column=4).value = "EUR 100"
    sheet.cell(row=2, column=5).value = "€1200"
    workbook.save(path)
    workbook.close()
    result = await ExtractBudgetRevisionEvidenceUseCase(
        parser=ExcelFileParser(), approved_scope_service=_ApprovedScopeService()
    ).execute(
        document=_document(), revision=_revision(path), source_path=path,
        existing_row_identities=frozenset(),
    )
    assert len(result.rows) == 1
    assert result.rows[0].unit_price == Decimal("100")
    assert result.rows[0].amount == Decimal("1200")



@pytest.mark.parametrize(
    ("column", "number_format", "currency"),
    [
        (4, '$#,##0.00', "EUR"),
        (5, '[$USD-409]#,##0.00', "EUR"),
        (4, '[$€-407]#,##0.00', "USD"),
        (5, '"£"#,##0.00', "EUR"),
        (4, '[$JPY-409]#,##0.00', "EUR"),
        (4, '¥#,##0.00', "EUR"),
        (5, '₹#,##0.00', "EUR"),
        (4, '₩#,##0.00', "EUR"),
        (5, '"¥"#,##0.00', "EUR"),
        (4, '€"₹"#,##0.00', "EUR"),
    ],
)
@pytest.mark.asyncio
async def test_budget_evidence_rejects_conflicting_numeric_cell_currency_formats(
    tmp_path: Path, column: int, number_format: str, currency: str,
) -> None:
    """A numeric Excel value can carry a conflicting currency in its number format."""
    path = tmp_path / "format-currency-conflict.xlsx"
    _budget_xlsx(path)
    workbook = openpyxl.load_workbook(path)
    workbook.active.cell(row=2, column=column).number_format = number_format
    workbook.active.cell(row=2, column=6).value = currency
    workbook.save(path)
    workbook.close()
    with pytest.raises(ExcelParsingError, match="currency (marker conflicts|format marker)"):
        await ExtractBudgetRevisionEvidenceUseCase(
            parser=ExcelFileParser(), approved_scope_service=_ApprovedScopeService(),
        ).execute(
            document=_document(), revision=_revision(path), source_path=path,
            existing_row_identities=frozenset(),
        )


@pytest.mark.parametrize(
    ("number_format", "currency"),
    [
        ('€#,##0.00', "EUR"),
        ('[$€-407]#,##0.00', "EUR"),
        ('[$-409]#,##0.00', "EUR"),
        ('[$-409]#,##0.00', "USD"),
    ],
)
@pytest.mark.asyncio
async def test_budget_evidence_accepts_matching_or_locale_only_currency_formats(
    tmp_path: Path, number_format: str, currency: str,
) -> None:
    """A locale directive alone never becomes a dollar denomination."""
    path = tmp_path / "format-currency-valid.xlsx"
    _budget_xlsx(path)
    workbook = openpyxl.load_workbook(path)
    sheet = workbook.active
    sheet.cell(row=2, column=4).number_format = number_format
    sheet.cell(row=2, column=6).value = currency
    workbook.save(path)
    workbook.close()
    result = await ExtractBudgetRevisionEvidenceUseCase(
        parser=ExcelFileParser(), approved_scope_service=_ApprovedScopeService(),
    ).execute(
        document=_document(), revision=_revision(path), source_path=path,
        existing_row_identities=frozenset(),
    )
    assert len(result.rows) == 1
    assert result.rows[0].unit_price == Decimal("100")
    assert result.rows[0].currency == currency



@pytest.mark.parametrize(
    "number_format",
    [
        '€#,##0.00;€-#,##0.00;€0.00;"₹"@',
        '€"#;#"0.00;€-#,##0.00;€0.00;"¥"@',
        r'€\;0.00;€-#,##0.00;€0.00;"₹"@',
    ],
)
@pytest.mark.asyncio
async def test_budget_numeric_evidence_ignores_text_only_format_section(
    tmp_path: Path, number_format: str,
) -> None:
    """The fourth Excel format section applies to text, not numeric amounts."""
    path = tmp_path / "numeric-with-text-format.xlsx"
    _budget_xlsx(path)
    workbook = openpyxl.load_workbook(path)
    workbook.active.cell(row=2, column=4).number_format = number_format
    workbook.save(path)
    workbook.close()

    result = await ExtractBudgetRevisionEvidenceUseCase(
        parser=ExcelFileParser(), approved_scope_service=_ApprovedScopeService(),
    ).execute(
        document=_document(), revision=_revision(path), source_path=path,
        existing_row_identities=frozenset(),
    )
    assert len(result.rows) == 1
    assert result.rows[0].currency == "EUR"
    assert result.rows[0].unit_price == Decimal("100")


@pytest.mark.asyncio
async def test_budget_currency_rejects_unsupported_symbol_in_numeric_format_section(
    tmp_path: Path,
) -> None:
    path = tmp_path / "numeric-section-unsupported.xlsx"
    _budget_xlsx(path)
    workbook = openpyxl.load_workbook(path)
    workbook.active.cell(row=2, column=4).number_format = '₹#,##0.00;€-#,##0.00;€0.00;"¥"@'
    workbook.save(path)
    workbook.close()

    with pytest.raises(ExcelParsingError, match="Unsupported budget currency format"):
        await ExtractBudgetRevisionEvidenceUseCase(
            parser=ExcelFileParser(), approved_scope_service=_ApprovedScopeService(),
        ).execute(
            document=_document(), revision=_revision(path), source_path=path,
            existing_row_identities=frozenset(),
        )


@pytest.mark.asyncio
async def test_budget_evidence_rejects_duplicate_canonical_headers(
    tmp_path: Path,
) -> None:
    path = tmp_path / "duplicate-total.xlsx"
    _budget_xlsx(path)
    wb = openpyxl.load_workbook(path)
    wb.active.cell(row=1, column=9).value = "Total (Gross)"
    wb.active.cell(row=2, column=9).value = 99999
    wb.save(path)
    wb.close()
    with pytest.raises(ExcelParsingError, match="Duplicate canonical budget evidence columns"):
        await ExcelFileParser().parse_budget(path, include_evidence=True)


@pytest.mark.asyncio
async def test_redelivery_of_retired_wbs_row_skips_new_authority_check(
    tmp_path: Path,
) -> None:
    """Previously ingested immutable rows need not relink a retired WBS node."""
    source = tmp_path / "budget.xlsx"
    _budget_xlsx(source)

    result = await ExtractBudgetRevisionEvidenceUseCase(
        parser=ExcelFileParser(), approved_scope_service=_DeniedApprovedScopeService()
    ).execute(
        document=_document(), revision=_revision(source), source_path=source,
        existing_row_identities=frozenset({(REVISION, "Budget 2026!2")}),
    )
    assert result.rows == ()
    assert result.assessment_reason == "no_verified_budget_rows"



@pytest.mark.asyncio
async def test_budget_historical_xlsx_revision_survives_current_bc3_format(
    tmp_path: Path,
) -> None:
    """Format changes in mutable document metadata do not rewrite source history."""
    source = tmp_path / "historical-budget.xlsx"
    _budget_xlsx(source)
    current = replace(_document(), file_format=".bc3", filename="budget.bc3")
    result = await ExtractBudgetRevisionEvidenceUseCase(
        parser=ExcelFileParser(), approved_scope_service=_ApprovedScopeService(),
    ).execute(
        document=current, revision=_revision(source), source_path=source,
        existing_row_identities=frozenset(),
    )
    assert len(result.rows) == 1
    assert result.rows[0].revision_id == REVISION


@pytest.mark.asyncio
async def test_budget_rejects_non_xlsx_selected_revision_with_valid_xlsx_bytes(
    tmp_path: Path,
) -> None:
    """Current XLSX metadata or an XLSX-shaped blob cannot attest a BC3 revision."""
    source = tmp_path / "budget.xlsx"
    _budget_xlsx(source)
    non_xlsx_revision = _revision(source).model_copy(
        update={"blob_key": "revisions/budget.bc3"}
    )
    with pytest.raises(InvalidBudgetRowEvidenceError, match="selected source revision"):
        await ExtractBudgetRevisionEvidenceUseCase(
            parser=ExcelFileParser(), approved_scope_service=_ApprovedScopeService(),
        ).execute(
            document=_document(), revision=non_xlsx_revision, source_path=source,
            existing_row_identities=frozenset(),
        )


@pytest.mark.asyncio
async def test_rejects_tampered_xlsx_before_parsing(tmp_path: Path) -> None:
    source = tmp_path / "budget.xlsx"
    _budget_xlsx(source)
    revision = _revision(source)
    source.write_bytes(source.read_bytes() + b"tampered")
    with pytest.raises(Exception, match="blob hash mismatch"):
        await ExtractBudgetRevisionEvidenceUseCase(parser=ExcelFileParser(), approved_scope_service=_ApprovedScopeService()).execute(document=_document(), revision=revision, source_path=source, existing_row_identities=frozenset())


@pytest.mark.asyncio
async def test_real_rows_with_mixed_units_or_monetary_bases_stay_not_assessed(
    tmp_path: Path,
) -> None:
    """Parser output must never become a green result when row semantics conflict."""
    source = tmp_path / "mixed-budget.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(["Item", "Quantity", "Unit", "Unit Price", "Total", "Currency", "Tax Basis"])
    worksheet.append(["Concrete", 12, "m3", 100, 1200, "EUR", "NET"])
    worksheet.append(["Rebar", 10, "kg", 90, 900, "USD", "GROSS"])
    workbook.save(source)

    result = await ExtractBudgetRevisionEvidenceUseCase(
        parser=ExcelFileParser(), approved_scope_service=_ApprovedScopeService()
    ).execute(
        document=_document(),
        revision=_revision(source),
        source_path=source,
        existing_row_identities=frozenset(),
    )

    assert len(result.rows) == 2
    assert result.assessment_reason == "incompatible_budget_row_semantics"


@pytest.mark.asyncio
async def test_redelivery_skips_an_already_recorded_revision_locator(tmp_path: Path) -> None:
    """The same immutable revision may be delivered again without duplicate evidence."""
    source = tmp_path / "budget.xlsx"
    _budget_xlsx(source)

    result = await ExtractBudgetRevisionEvidenceUseCase(
        parser=ExcelFileParser(), approved_scope_service=_ApprovedScopeService()
    ).execute(
        document=_document(),
        revision=_revision(source),
        source_path=source,
        existing_row_identities=frozenset({(REVISION, "Budget 2026!2")}),
    )

    assert result.rows == ()
    assert result.assessment_reason == "no_verified_budget_rows"
@pytest.mark.parametrize(
    ("field_name", "invalid_value"),
    [
        ("document_type", DocumentType.CONTRACT),
        ("id", UUID("77777777-7777-7777-7777-777777777777")),
        ("tenant_id", UUID("88888888-8888-8888-8888-888888888888")),
        ("project_id", UUID("99999999-9999-9999-9999-999999999999")),
    ],
)
@pytest.mark.asyncio
async def test_rejects_wrong_budget_document_identity_before_parsing(
    tmp_path: Path, field_name: str, invalid_value: object
) -> None:
    """Cross-document, cross-tenant and cross-project revision substitution fail closed."""
    source = tmp_path / "budget.xlsx"
    _budget_xlsx(source)
    mismatched_document = replace(_document(), **{field_name: invalid_value})

    with pytest.raises(InvalidBudgetRowEvidenceError):
        await ExtractBudgetRevisionEvidenceUseCase(
            parser=ExcelFileParser(), approved_scope_service=_ApprovedScopeService()
        ).execute(
            document=mismatched_document,
            revision=_revision(source),
            source_path=source,
            existing_row_identities=frozenset(),
        )


class _DeniedApprovedScopeService:
    async def require(
        self, *, tenant_id: UUID, project_id: UUID, node_ids: set[UUID]
    ) -> UUID:
        """Represent the authoritative WBS gate denying an unapproved node."""
        raise UnapprovedWBSMappingError("WBS baseline/node not approved")


@pytest.mark.asyncio
async def test_denied_wbs_scope_is_propagated_without_emitting_budget_rows(
    tmp_path: Path,
) -> None:
    """A denied WBS mapping must not be downgraded into an unscoped budget row."""
    source = tmp_path / "budget.xlsx"
    _budget_xlsx(source)

    with pytest.raises(UnapprovedWBSMappingError, match="not approved"):
        await ExtractBudgetRevisionEvidenceUseCase(
            parser=ExcelFileParser(), approved_scope_service=_DeniedApprovedScopeService()
        ).execute(
            document=_document(),
            revision=_revision(source),
            source_path=source,
            existing_row_identities=frozenset(),
        )
