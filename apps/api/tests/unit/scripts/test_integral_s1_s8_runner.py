"""Unit guards for the integral S1-S8 UAT: its independent PostgreSQL judge, the manifest, the identity preflight,
the tenancy seed and the committed fixtures the browser run is pinned to.

The judge is a PURE function over rows, so every claim is proved here both ways: the honest run passes, and each
tampering (a shared tenant, a self-approval nobody declared, an approval of the other revision, a rejected draft that
became trusted, a link on another baseline, snapshots that differ after the relogin ...) fails exactly its check.
"""

from __future__ import annotations

import asyncio
import copy
import hashlib
import json
import os
from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4

import pytest

from scripts import preflight_integral_identities as preflight
from scripts import seed_integral_uat as seed
from scripts import verify_integral_s1_s8_pg as verifier

WEB = Path(__file__).resolve().parents[4] / "web"
REPO_ROOT = Path(__file__).resolve().parents[5]
FIXTURES = WEB / "src" / "tests" / "e2e" / "test-data" / "triplet"
TRIPLET = json.loads((FIXTURES / "triplet.manifest.json").read_text())
EXTRA = json.loads((FIXTURES / "integral.manifest.json").read_text())

REASON = "The draft sets an implausible strength."
SHA = {
    "contractA": TRIPLET["sha256"]["triplet-contract.docx"],
    "contractB": EXTRA["contract_b"]["sha256"],
    "budget": TRIPLET["sha256"]["triplet-budget.xlsx"],
    "schedule": TRIPLET["sha256"]["triplet-schedule.xlsx"],
    "specification": TRIPLET["sha256"]["triplet-specification.docx"],
    "specificationRejected": EXTRA["specification_rejected"]["sha256"],
}


_LAST: list = []


def _world() -> tuple[verifier.Facts, dict, verifier.Expectations]:
    """An honest, complete run: every row the journey leaves, as PostgreSQL would hold it."""
    tenant_a, tenant_b = uuid4(), uuid4()
    project_a, project_b = uuid4(), uuid4()
    user_a, user_b = uuid4(), uuid4()
    docs = {
        k: uuid4() for k in ("contract", "budget", "schedule", "specification", "rejectedCandidate")
    }
    types = {
        "contract": "contract",
        "budget": "budget",
        "schedule": "schedule",
        "specification": "technical_spec",
        "rejectedCandidate": "technical_spec",
    }
    shas = {
        "contract": None,
        "budget": SHA["budget"],
        "schedule": SHA["schedule"],
        "specification": SHA["specification"],
        "rejectedCandidate": SHA["specificationRejected"],
    }
    rev_a, rev_b = uuid4(), uuid4()
    revisions = [
        {
            "document_id": docs["contract"],
            "revision_id": rev_a,
            "rev_no": 1,
            "blob_hash": SHA["contractA"],
        },
        {
            "document_id": docs["contract"],
            "revision_id": rev_b,
            "rev_no": 2,
            "blob_hash": SHA["contractB"],
        },
    ]
    revision_of = {}
    for key in ("budget", "schedule", "specification", "rejectedCandidate"):
        revision_of[key] = uuid4()
        revisions.append(
            {
                "document_id": docs[key],
                "revision_id": revision_of[key],
                "rev_no": 1,
                "blob_hash": shas[key],
            }
        )
    art = {k: uuid4() for k in ("a", "b", "spec", "rej")}
    artifacts = [
        {
            "artifact_id": art["a"],
            "document_id": docs["contract"],
            "document_revision_id": rev_a,
            "artifact_version": 1,
            "artifact_hash": "h-a",
            "trust_state": "trusted",
            "lifecycle_status": "superseded",
        },
        {
            "artifact_id": art["b"],
            "document_id": docs["contract"],
            "document_revision_id": rev_b,
            "artifact_version": 2,
            "artifact_hash": "h-b",
            "trust_state": "trusted",
            "lifecycle_status": "active",
        },
        {
            "artifact_id": art["spec"],
            "document_id": docs["specification"],
            "document_revision_id": revision_of["specification"],
            "artifact_version": 1,
            "artifact_hash": "h-s",
            "trust_state": "trusted",
            "lifecycle_status": "active",
        },
        {
            "artifact_id": art["rej"],
            "document_id": docs["rejectedCandidate"],
            "document_revision_id": revision_of["rejectedCandidate"],
            "artifact_version": 1,
            "artifact_hash": "h-r",
            "trust_state": "rejected",
            "lifecycle_status": "active",
        },
    ]

    def review(doc: str, status: str, artifact: str, artifact_hash: str, **meta: str) -> dict:
        return {
            "project_id": project_a,
            "document_id": docs[doc],
            "item_id": uuid4(),
            "id": uuid4(),
            "status": status,
            "approved_by": "reviewer",
            "approved_at": "2026-10-31T10:00:00",
            "review_metadata": {
                "candidate_binding": {"artifact_id": artifact, "artifact_hash": artifact_hash},
                **meta,
            },
        }

    reviews = [
        review("contract", "APPROVED", str(art["a"]), "h-a"),
        review("contract", "APPROVED", str(art["b"]), "h-b"),
        review("specification", "APPROVED", str(art["spec"]), "h-s"),
        review("rejectedCandidate", "REJECTED", str(art["rej"]), "h-r", rejection_reason=REASON),
    ]
    clause = uuid4()
    baseline, change_set, node = uuid4(), uuid4(), uuid4()
    events = [
        {
            "event_id": uuid4(),
            "event_type": "graph.completed",
            "actor": "system",
            "payload": {},
            "document_id": None,
        },
        {
            "event_id": uuid4(),
            "event_type": "revision.changed",
            "actor": "system",
            "payload": {},
            "document_id": str(docs["contract"]),
        },
        {
            "event_id": uuid4(),
            "event_type": "budget.wbs_link_approved",
            "actor": f"user:{user_a}",
            "document_id": None,
            "payload": {
                "source_revision_id": str(revision_of["budget"]),
                "source_blob_hash": SHA["budget"],
                "source_locator": "Presupuesto!2",
                "wbs_baseline_id": str(baseline),
                "wbs_node_id": str(node),
            },
        },
        {
            "event_id": uuid4(),
            "event_type": "schedule.wbs_link_approved",
            "actor": f"user:{user_a}",
            "document_id": None,
            "payload": {
                "source_revision_id": str(revision_of["schedule"]),
                "source_blob_hash": SHA["schedule"],
                "source_locator": "Plan!4",
                "wbs_baseline_id": str(baseline),
                "wbs_node_id": str(node),
            },
        },
        {
            "event_id": uuid4(),
            "event_type": "technical.wbs_link_approved",
            "actor": f"user:{user_a}",
            "document_id": None,
            "payload": {
                "technical_revision_id": str(revision_of["specification"]),
                "source_blob_hash": SHA["specification"],
                "source_section": "4.1.2",
                "wbs_baseline_id": str(baseline),
                "wbs_node_id": str(node),
            },
        },
    ]
    facts = verifier.Facts(
        project_a={"id": project_a, "tenant_id": tenant_a, "name": "A"},
        project_b={"id": project_b, "tenant_id": tenant_b, "name": "B"},
        tenants={
            str(tenant_a): {"id": tenant_a, "clerk_org_id": "org_a"},
            str(tenant_b): {"id": tenant_b, "clerk_org_id": "org_b"},
        },
        users=[
            {"id": user_a, "tenant_id": tenant_a, "clerk_user_id": "user_a", "is_active": True},
            {"id": user_b, "tenant_id": tenant_b, "clerk_user_id": "user_b", "is_active": True},
        ],
        documents=[
            {
                "id": docs[k],
                "tenant_id": tenant_a,
                "project_id": project_a,
                "document_type": types[k],
            }
            for k in docs
        ],
        revisions=revisions,
        artifacts=artifacts,
        reviews=reviews,
        alerts=[
            {
                "project_id": project_a,
                "tenant_id": tenant_a,
                "rule_code": rule,
                "status": "open",
                "source_clause_id": clause,
            }
            for rule in TRIPLET["expected_rules"]
        ],
        clause_revisions={str(clause): str(rev_a)},
        events=events,
        graph_events=1,
        analyses=1,
        six_dimension_snapshots=1,
        baselines=[
            {
                "id": baseline,
                "project_id": project_a,
                "baseline_no": 1,
                "approved_by": user_a,
                "source_change_set_id": change_set,
            }
        ],
        change_sets=[
            {
                "id": change_set,
                "submitted_by": user_a,
                "decided_by": user_a,
                "decided_self_approval": True,
            }
        ],
        tenant_a_distinct_approver_required=False,
        tenant_a_ai_spend=0.42,
    )
    browser = {
        "mode": "real-ai-uat",
        "evidenceClass": "UAT_REAL",
        "projectIdA": str(project_a),
        "projectIdB": str(project_b),
        "documentIds": {k: str(v) for k, v in docs.items()},
        "clerkIdentityChecks": {
            "distinctUsers": True,
            "distinctOrganizations": True,
            "distinctTenants": True,
        },
        "snapshotDigests": {"before": "d" * 64, "after": "d" * 64},
        "stages": [
            {"id": f"S{i}", "status": "PASS", "partial": [], "evidence": {}} for i in range(1, 9)
        ],
    }
    browser["stages"][4]["evidence"] = {
        "rejection": {"reasonSha256": hashlib.sha256(REASON.encode()).hexdigest()}
    }
    browser["stages"][6]["evidence"] = {"wbs": {"selfApproved": True}}
    expect = verifier.Expectations(
        fixture_sha=SHA,
        expected_rules=list(TRIPLET["expected_rules"]),
        clerk_user_a="user_a",
        clerk_org_a="org_a",
        clerk_user_b="user_b",
        clerk_org_b="org_b",
        tenant_a=str(tenant_a),
        tenant_b=str(tenant_b),
        real_ai=True,
        worker_provider_successes=3,
        reason_sha256=hashlib.sha256(REASON.encode()).hexdigest(),
        self_approved_declared=True,
    )
    _LAST[:] = [facts, browser, expect]
    return facts, browser, expect


def _failed(facts: verifier.Facts, browser: dict, expect: verifier.Expectations) -> list[str]:
    return [f"{c.stage}:{c.name}" for c in verifier.judge(facts, browser, expect) if not c.ok]


def test_an_honest_run_passes_every_postgres_check() -> None:
    facts, browser, expect = _world()
    assert _failed(facts, browser, expect) == []


def test_two_projects_in_one_tenant_fail_isolation() -> None:
    facts, browser, expect = _world()
    facts.project_b["tenant_id"] = facts.project_a["tenant_id"]
    failed = _failed(facts, browser, expect)
    assert "S1:the two projects belong to two different tenants" in failed


def test_a_tenant_not_bound_to_the_clerk_organization_fails() -> None:
    facts, browser, expect = _world()
    facts.tenants[str(facts.project_b["tenant_id"])]["clerk_org_id"] = "org_a"
    assert "S1:tenant B is bound to the Clerk Organization of identity B" in _failed(
        facts, browser, expect
    )


def test_a_missing_second_identity_cannot_pass() -> None:
    facts, browser, expect = _world()
    expect.clerk_user_b = ""
    assert "S1:identity B is an active user of tenant B" in _failed(facts, browser, expect)


def test_changed_stored_bytes_fail_the_pin() -> None:
    facts, browser, expect = _world()
    facts.revisions[0]["blob_hash"] = "0" * 64
    assert any(
        name.startswith("S2:contract: stored revisions") for name in _failed(facts, browser, expect)
    )


def test_a_missing_document_fails() -> None:
    facts, browser, expect = _world()
    facts.documents.pop()
    assert "S2:project A holds exactly five documents" in _failed(facts, browser, expect)


def test_no_persisted_alert_fails_s3_and_s4() -> None:
    facts, browser, expect = _world()
    facts.alerts = []
    failed = _failed(facts, browser, expect)
    assert "S3:alerts are persisted for project A, all of tenant A" in failed
    assert "S4:an alert cites a clause" in failed


def test_an_alert_clause_outside_the_contract_fails_s4() -> None:
    facts, browser, expect = _world()
    facts.clause_revisions = {key: str(uuid4()) for key in facts.clause_revisions}
    assert "S4:the clause belongs to a revision of the contract" in _failed(facts, browser, expect)


def test_real_mode_demands_a_provider_call_and_recorded_spend() -> None:
    facts, browser, expect = _world()
    expect.worker_provider_successes = 0
    facts.tenant_a_ai_spend = 0
    failed = _failed(facts, browser, expect)
    assert "S3:the worker made at least one real provider call" in failed
    assert "S3:the tenant's AI spend was recorded" in failed


def test_a_rehearsal_is_not_held_to_provider_calls() -> None:
    facts, browser, expect = _world()
    expect.real_ai = False
    expect.worker_provider_successes = 0
    facts.tenant_a_ai_spend = 0
    assert _failed(facts, browser, expect) == []


def test_the_approval_of_revision_a_never_satisfies_revision_b() -> None:
    facts, browser, expect = _world()
    facts.reviews = [
        r
        for r in facts.reviews
        if (r["review_metadata"].get("candidate_binding") or {}).get("artifact_hash") != "h-b"
    ]
    assert any(
        name.startswith("S5:contract: every trusted artifact has an approved review")
        for name in _failed(facts, browser, expect)
    )


def test_an_approval_that_names_nobody_fails() -> None:
    facts, browser, expect = _world()
    facts.reviews[2]["approved_by"] = ""
    assert any(
        name.startswith("S5:specification: every trusted artifact")
        for name in _failed(facts, browser, expect)
    )


def test_a_rejected_draft_that_became_trusted_fails() -> None:
    facts, browser, expect = _world()
    facts.artifacts[3]["trust_state"] = "trusted"
    assert "S5:nothing of the rejected draft is trusted" in _failed(facts, browser, expect)


def test_a_different_rejection_reason_fails() -> None:
    facts, browser, expect = _world()
    facts.reviews[3]["review_metadata"]["rejection_reason"] = "another reason"
    assert "S5:the stored rejection reason is the typed one" in _failed(facts, browser, expect)


def test_a_pending_review_left_behind_fails() -> None:
    facts, browser, expect = _world()
    facts.reviews[0]["status"] = "PENDING_REVIEW_REQUIRED"
    assert "S5:no review is left pending" in _failed(facts, browser, expect)


def test_revision_b_must_be_the_active_trusted_artifact() -> None:
    facts, browser, expect = _world()
    facts.artifacts[1]["trust_state"] = "proposed"
    failed = _failed(facts, browser, expect)
    assert "S6:revision B is the active trusted artifact" in failed
    assert "S6:no artifact of the project is left proposed" in failed


def test_a_duplicated_change_event_fails() -> None:
    facts, browser, expect = _world()
    facts.events.append(dict(facts.events[1], event_id=uuid4()))
    assert "S6:exactly one revision.changed exists for the contract" in _failed(
        facts, browser, expect
    )


def test_a_self_approval_nobody_declared_fails() -> None:
    facts, browser, expect = _world()
    facts.tenant_a_distinct_approver_required = True
    assert "S7:a self-approval happened only because tenant A declares it" in _failed(
        facts, browser, expect
    )


def test_the_browser_and_the_database_must_agree_on_self_approval() -> None:
    facts, browser, expect = _world()
    expect.self_approved_declared = False
    assert "S7:the self-approval flag equals what the browser recorded" in _failed(
        facts, browser, expect
    )


def test_a_missing_family_link_fails() -> None:
    facts, browser, expect = _world()
    facts.events = [e for e in facts.events if e["event_type"] != "budget.wbs_link_approved"]
    assert "S7:exactly one approved budget link exists" in _failed(facts, browser, expect)


def test_a_link_on_another_baseline_fails() -> None:
    facts, browser, expect = _world()
    facts.events[2]["payload"] = dict(facts.events[2]["payload"], wbs_baseline_id=str(uuid4()))
    assert "S7:budget link: it names the approved baseline and node" in _failed(
        facts, browser, expect
    )


def test_a_link_with_the_wrong_blob_hash_fails() -> None:
    facts, browser, expect = _world()
    facts.events[3]["payload"] = dict(facts.events[3]["payload"], source_blob_hash="f" * 64)
    assert "S7:schedule link: it names a stored source revision and its blob hash" in _failed(
        facts, browser, expect
    )


def test_a_link_by_a_non_user_actor_fails() -> None:
    facts, browser, expect = _world()
    facts.events[4]["actor"] = "system"
    assert "S7:technical link: actor is a human user of tenant A" in _failed(facts, browser, expect)


def test_snapshots_that_differ_after_the_relogin_fail_s8() -> None:
    facts, browser, expect = _world()
    browser["snapshotDigests"] = {"before": "a" * 64, "after": "b" * 64}
    assert "S8:the browser's snapshots before and after the new login are identical" in _failed(
        facts, browser, expect
    )


def test_a_missing_project_stops_the_judgement_early() -> None:
    facts, browser, expect = _world()
    facts.project_b = None
    assert _failed(facts, browser, expect) == ["S1:project B exists in PostgreSQL"]


# ── manifest ────────────────────────────────────────────────────────────────────────────────────────────────
def _manifest(browser: dict, checks: list[verifier.Check] | None = None) -> dict:
    facts, _, expect = _LAST
    return verifier.build_manifest(
        browser,
        checks if checks is not None else verifier.judge(facts, browser, expect),
        {"worker_provider_successes": 2},
        {"candidate_sha": "c" * 40},
    )


def test_a_complete_real_run_is_complete_but_never_self_accepted() -> None:
    _, browser, _ = _world()
    manifest = _manifest(browser)
    assert manifest["verdict"] == "UAT_REAL_COMPLETE"
    assert manifest["acceptance"] == "NOT_ACCEPTED_UNTIL_INDEPENDENT_REVIEW"
    assert (
        manifest["project_id_A"] == browser["projectIdA"]
        and manifest["project_id_B"] == browser["projectIdB"]
    )
    assert manifest["candidate_sha"] == "c" * 40
    assert set(manifest["steps_S1_S8"]) == set(verifier.STAGES)


def test_declared_gaps_downgrade_the_verdict() -> None:
    _, browser, _ = _world()
    browser["stages"][4]["status"] = "PARTIAL"
    browser["stages"][4]["partial"] = ["granular decisions have no route"]
    manifest = _manifest(browser)
    assert manifest["verdict"] == "UAT_REAL_WITH_DECLARED_GAPS"
    assert manifest["steps_S1_S8"]["S5"]["declaredGaps"] == ["granular decisions have no route"]


def test_a_deterministic_rehearsal_is_never_uat() -> None:
    _, browser, _ = _world()
    browser["mode"] = "deterministic-rehearsal"
    assert _manifest(browser)["verdict"] == "REHEARSAL_ONLY_NOT_UAT"


def test_the_first_failure_is_the_browsers_and_later_stages_stay_not_run() -> None:
    _, browser, _ = _world()
    browser["stages"][6]["status"] = "FAIL"
    for index in (7,):
        browser["stages"][index]["status"] = "NOT_RUN"
    browser["firstFailure"] = {
        "stage": "S7",
        "message": "budget: HTTP 400 source revision is not trusted",
    }
    manifest = _manifest(browser, [])
    assert manifest["verdict"] == "FAILED"
    assert manifest["first_failure"]["stage"] == "S7"
    assert manifest["steps_S1_S8"]["S8"]["verdict"] == "NOT_RUN"


def test_a_failed_postgres_check_fails_the_stage_even_when_the_browser_passed() -> None:
    facts, browser, expect = _world()
    facts.artifacts[3]["trust_state"] = "trusted"
    manifest = _manifest(browser, verifier.judge(facts, browser, expect))
    assert manifest["steps_S1_S8"]["S5"]["verdict"] == "FAIL"
    assert manifest["steps_S1_S8"]["S5"]["failedChecks"]
    assert manifest["verdict"] == "FAILED"


def test_stages_nobody_checked_are_incomplete_not_passed() -> None:
    _, browser, _ = _world()
    manifest = _manifest(browser, [])
    assert manifest["verdict"] == "INCOMPLETE"


def test_the_manifest_carries_no_secret_shaped_value() -> None:
    _, browser, _ = _world()
    text = json.dumps(_manifest(browser))
    for forbidden in ("Bearer ", "sk-", "eyJ", "cookie", "password"):
        assert forbidden not in text


# ── collection against a fake connection (no database) ─────────────────────────────────────────────────────
class _Result:
    def __init__(self, rows: list[dict]) -> None:
        self._rows = rows

    def all(self) -> list[SimpleNamespace]:
        return [SimpleNamespace(_mapping=row) for row in self._rows]

    def scalar_one(self) -> int:
        return int(next(iter(self._rows[0].values()))) if self._rows else 0


class _FakeConnection:
    def __init__(self, answers: dict[str, list[dict]]) -> None:
        self.answers = answers
        self.statements: list[str] = []

    async def execute(self, statement: object, params: dict | None = None) -> _Result:
        sql = str(statement)
        self.statements.append(sql)
        assert (
            not sql.lstrip()
            .upper()
            .startswith(("INSERT", "UPDATE", "DELETE", "DROP", "ALTER", "TRUNCATE"))
        ), "the verifier must stay read-only"
        for fragment, rows in self.answers.items():
            if fragment in sql:
                return _Result(rows)
        return _Result([])


def test_collect_reads_only_and_assembles_the_facts() -> None:
    project_a, project_b, tenant = uuid4(), uuid4(), uuid4()
    clause, rev = uuid4(), uuid4()
    conn = _FakeConnection(
        {
            "FROM projects": [
                {"id": project_a, "tenant_id": tenant, "name": "A"},
                {"id": project_b, "tenant_id": uuid4(), "name": "B"},
            ],
            "FROM tenants": [
                {
                    "id": tenant,
                    "clerk_org_id": "org",
                    "settings": {"wbs_governance": {"require_distinct_approver": False}},
                    "ai_spend_current": 1.5,
                }
            ],
            "FROM alerts": [
                {
                    "project_id": project_a,
                    "tenant_id": tenant,
                    "rule_code": "R",
                    "status": "open",
                    "source_clause_id": clause,
                }
            ],
            "FROM clauses": [{"id": clause, "revision_id": rev}],
            "FROM project_events": [
                {
                    "event_id": uuid4(),
                    "event_type": "graph.completed",
                    "actor": "x",
                    "payload": {},
                    "source_revision_id": None,
                    "document_id": None,
                }
            ],
            "FROM analyses": [{"n": 2}],
            "FROM project_snapshots": [{"n": 1}],
        }
    )
    facts = asyncio.run(verifier.collect(conn, project_a, project_b))
    assert facts.project_a["id"] == project_a
    assert facts.project_b["id"] == project_b
    assert facts.clause_revisions == {str(clause): str(rev)}
    assert (facts.graph_events, facts.analyses, facts.six_dimension_snapshots) == (1, 2, 1)
    assert facts.tenant_a_distinct_approver_required is False
    assert facts.tenant_a_ai_spend == 1.5
    assert conn.statements, "something was read"


def test_collect_stops_when_project_a_is_missing() -> None:
    conn = _FakeConnection({"FROM projects": []})
    facts = asyncio.run(verifier.collect(conn, uuid4(), uuid4()))
    assert facts.project_a is None
    assert facts.documents == []


# ── CLI hygiene ─────────────────────────────────────────────────────────────────────────────────────────────
def test_the_evidence_files_have_fixed_places_inside_the_checkout() -> None:
    """No path comes from the command line: every read and write is at a fixed place of this checkout."""
    for place in (
        verifier.BROWSER_MANIFEST,
        verifier.PROVIDER_RECEIPT,
        verifier.MANIFEST_OUTPUT,
        verifier.FIXTURE_DIR,
    ):
        assert place.is_relative_to(verifier.REPO_ROOT)
    source = (Path(verifier.__file__)).read_text()
    assert "--output" not in source
    assert "--browser-manifest" not in source


def test_database_urls_are_normalised_for_asyncpg() -> None:
    assert verifier._normalize_database_url("postgresql://u@h/d") == "postgresql+asyncpg://u@h/d"
    assert verifier._normalize_database_url("postgres://u@h/d") == "postgresql+asyncpg://u@h/d"
    assert (
        verifier._normalize_database_url("postgresql+asyncpg://u@h/d")
        == "postgresql+asyncpg://u@h/d"
    )


# ── committed fixtures ──────────────────────────────────────────────────────────────────────────────────────
def test_every_committed_fixture_has_its_pinned_sha() -> None:
    files = {
        "triplet-contract.docx": TRIPLET["sha256"]["triplet-contract.docx"],
        "triplet-budget.xlsx": TRIPLET["sha256"]["triplet-budget.xlsx"],
        "triplet-schedule.xlsx": TRIPLET["sha256"]["triplet-schedule.xlsx"],
        "triplet-specification.docx": TRIPLET["sha256"]["triplet-specification.docx"],
        "triplet-contract-b.docx": EXTRA["contract_b"]["sha256"],
        "triplet-specification-rejected.docx": EXTRA["specification_rejected"]["sha256"],
    }
    for name, pinned in files.items():
        assert hashlib.sha256((FIXTURES / name).read_bytes()).hexdigest() == pinned, name


def test_contract_b_adds_one_clause_and_changes_no_finding_fact() -> None:
    import docx

    a = [p.text for p in docx.Document(str(FIXTURES / "triplet-contract.docx")).paragraphs]
    b = [p.text for p in docx.Document(str(FIXTURES / "triplet-contract-b.docx")).paragraphs]
    assert b[: len(a)] == a, (
        "revision B keeps every clause of A (price and completion date unchanged)"
    )
    assert b[len(a) :] == [EXTRA["contract_b"]["added_sentence"]]


def test_the_runner_row_locators_are_the_ones_the_product_parser_names() -> None:
    """The browser links ``Presupuesto!2`` and ``Plan!4``: they must exist, once, in the pinned workbooks."""
    from src.documents.adapters.parsers.excel_file_parser import ExcelFileParser

    async def locators() -> tuple[list[str], list[str]]:
        parser = ExcelFileParser()
        budget = await parser.parse_budget(FIXTURES / "triplet-budget.xlsx", include_evidence=True)
        schedule = await parser.parse_schedule(
            FIXTURES / "triplet-schedule.xlsx", include_evidence=True
        )
        return [str(r["source_locator"]) for r in budget], [
            str(r["source_locator"]) for r in schedule
        ]

    budget, schedule = asyncio.run(locators())
    assert budget.count("Presupuesto!2") == 1
    assert schedule.count("Plan!4") == 1
    source = (WEB / "src" / "tests" / "e2e" / "support" / "integral" / "fixtures.ts").read_text()
    assert 'budget: "Presupuesto!2"' in source
    assert 'schedule: "Plan!4"' in source


# ── identity preflight ──────────────────────────────────────────────────────────────────────────────────────
def _membership(org: str, tenant: str | None) -> dict:
    return {"organization": {"id": org, "public_metadata": {"tenant_id": tenant} if tenant else {}}}


def _tid(n: int) -> str:
    return f"00000000-0000-0000-0000-{n:012d}"


def _ready(label: str, user: str, org: str, tenant: str) -> preflight.IdentityReport:
    tenant = _tid(int(tenant[1:])) if tenant[:1] == "t" and tenant[1:].isdigit() else tenant
    return preflight.evaluate_identity(label, [{"id": user}], [_membership(org, tenant)])


def test_two_identities_in_two_tenants_are_ready() -> None:
    code, reasons = preflight.evaluate_pair(
        _ready("A", "u1", "o1", "t1"), _ready("B", "u2", "o2", "t2")
    )
    assert (code, reasons) == (preflight.READY, [])


def test_a_missing_b_user_fails_with_a_reason() -> None:
    b = preflight.evaluate_identity("B", [], [])
    code, reasons = preflight.evaluate_pair(_ready("A", "u1", "o1", "t1"), b)
    assert code == preflight.FAILED
    assert any("0 Clerk user" in r for r in reasons)


def test_a_b_user_without_an_organization_asks_for_provisioning() -> None:
    b = preflight.evaluate_identity("B", [{"id": "u2"}], [])
    code, _ = preflight.evaluate_pair(_ready("A", "u1", "o1", "t1"), b)
    assert code == preflight.NEEDS_B_ORGANIZATION


def test_an_ambiguous_email_fails() -> None:
    report = preflight.evaluate_identity("B", [{"id": "u2"}, {"id": "u3"}], [])
    assert not report.ready
    assert "exactly one" in report.problems[0]


def test_two_memberships_fail_because_the_tenant_would_be_ambiguous() -> None:
    report = preflight.evaluate_identity(
        "B", [{"id": "u2"}], [_membership("o2", "t2"), _membership("o3", "t3")]
    )
    assert not report.ready


def test_an_organization_without_a_tenant_fails() -> None:
    report = preflight.evaluate_identity("B", [{"id": "u2"}], [_membership("o2", None)])
    assert "tenant_id" in report.problems[0]


@pytest.mark.parametrize(
    ("b", "needle"),
    [
        (("u1", "o2", "t2"), "same Clerk user"),
        (("u2", "o1", "t2"), "share one Organization"),
        (("u2", "o2", "t1"), "share one tenant id"),
    ],
)
def test_shared_user_organization_or_tenant_fails(b: tuple[str, str, str], needle: str) -> None:
    code, reasons = preflight.evaluate_pair(_ready("A", "u1", "o1", "t1"), _ready("B", *b))
    assert code == preflight.FAILED
    assert any(needle in r for r in reasons)


def test_a_tenant_that_is_not_a_uuid_fails() -> None:
    report = preflight.evaluate_identity("B", [{"id": "u2"}], [_membership("o2", "not-a-uuid")])
    assert not report.ready
    assert "valid UUID" in report.problems[0]


def test_only_the_dedicated_development_b_variable_is_read() -> None:
    assert (
        preflight.resolve_b_email({"INTEGRAL_CLERK_DEV_B_EMAIL": " dev-b@example.test "})
        == "dev-b@example.test"
    )
    assert preflight.resolve_b_email({"E2E_CLERK_TENANT_B_USER_EMAIL": "prod-b@x.test"}) == ""
    assert preflight.resolve_b_email({}) == ""


def test_b_exports_come_from_the_report_and_name_the_seed_tenant() -> None:
    lines = dict(
        line.split("=", 1) for line in preflight.b_export_lines(_ready("B", "u2", "o2", "t2"))
    )
    assert (
        lines["TENANT_B_CLERK_E2E_ORG_ID"] == "o2" and lines["TENANT_B_CLERK_E2E_USER_ID"] == "u2"
    )
    assert lines["E2E_TENANT_B_ID"] == lines["TENANT_B_E2E_EXPECTED_TENANT_ID"] == _tid(2)


def test_exit_code_three_is_a_stop_not_a_repair() -> None:
    b = preflight.evaluate_identity("B", [{"id": "u2"}], [])
    code, _ = preflight.evaluate_pair(_ready("A", "u1", "o1", "t1"), b)
    assert code == preflight.NEEDS_B_ORGANIZATION


def _integral_job_text() -> str:
    text = (REPO_ROOT / ".github" / "workflows" / "acc01-s2-live-ai-uat.yml").read_text()
    return text[text.index("\n  integral-s1-s8:") :]


def test_the_uat_job_never_writes_to_clerk_b() -> None:
    job = _integral_job_text()
    assert "--export-github-env" in job
    assert job.count("provision_clerk_e2e_fixture.py") == 1  # identity A only
    assert "--export-prefix" not in job
    assert "E2E_CLERK_TENANT_B_USER_EMAIL: ${{ secrets.E2E_CLERK_TENANT_B_USER_EMAIL" not in job


def test_no_workflow_env_reads_the_legacy_b_secret() -> None:
    text = (REPO_ROOT / ".github" / "workflows" / "acc01-s2-live-ai-uat.yml").read_text()
    assert (
        "secrets.E2E_CLERK_TENANT_B_USER_EMAIL" not in text
        and "vars.E2E_CLERK_TENANT_B_USER_EMAIL" not in text
    )


def test_preparing_b_is_a_separate_sha_confirmed_job_that_is_not_a_uat() -> None:
    text = (REPO_ROOT / ".github" / "workflows" / "acc01-s2-live-ai-uat.yml").read_text()
    job = text[text.index("\n  integral-prepare-dev-b:") : text.index("\n  integral-s1-s8:")]
    assert (
        "integral_uat != 'true'" in job
        and "integral_confirm_prepare_dev_b_sha == github.sha" in job
    )
    assert "INTEGRAL_CLERK_DEV_B_EMAIL" in job


def test_main_fails_closed_without_b_and_never_reaches_clerk(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    for name in ("INTEGRAL_CLERK_DEV_B_EMAIL", "E2E_CLERK_TENANT_B_USER_EMAIL"):
        monkeypatch.delenv(name, raising=False)
    monkeypatch.setenv("CLERK_SECRET_KEY", "sk_test_x")
    assert preflight.main([]) == preflight.FAILED


def test_main_refuses_a_production_key_before_any_clerk_call(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("CLERK_SECRET_KEY", "sk_live_x")
    monkeypatch.setenv("INTEGRAL_CLERK_DEV_B_EMAIL", "b@example.test")
    monkeypatch.setenv("E2E_CLERK_USER_EMAIL", "a@example.test")
    assert preflight.main([]) == preflight.FAILED


def test_main_refuses_identical_a_and_b_emails(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("CLERK_SECRET_KEY", "sk_test_x")
    monkeypatch.setenv("E2E_CLERK_USER_EMAIL", "Same@example.test")
    monkeypatch.setenv("INTEGRAL_CLERK_DEV_B_EMAIL", "same@example.test")
    assert preflight.main([]) == preflight.FAILED


def test_a_broken_a_fails_before_judging_b() -> None:
    a = preflight.evaluate_identity("A", [], [])
    code, _ = preflight.evaluate_pair(a, _ready("B", "u2", "o2", "t2"))
    assert code == preflight.FAILED


def test_the_public_report_holds_no_clerk_identifier() -> None:
    report = _ready("B", "user_secret_id", "org_secret_id", "tenant")
    text = json.dumps(report.public())
    assert "user_secret_id" not in text
    assert "org_secret_id" not in text


# ── tenancy seed preconditions ──────────────────────────────────────────────────────────────────────────────
GOOD_ENV = {
    "ENVIRONMENT": "test",
    "STORAGE_PROVIDER": "local",
    "CLERK_E2E_USER_ID": "user_a",
    "CLERK_E2E_ORG_ID": "org_a",
    "TENANT_B_CLERK_E2E_USER_ID": "user_b",
    "TENANT_B_CLERK_E2E_ORG_ID": "org_b",
    "E2E_CLERK_TENANT_B_USER_EMAIL": "b@example.test",
}


def test_the_seed_accepts_two_distinct_identities() -> None:
    assert seed.require_environment(GOOD_ENV)["TENANT_B_CLERK_E2E_ORG_ID"] == "org_b"


@pytest.mark.parametrize(
    ("override", "message"),
    [
        ({"ENVIRONMENT": "production"}, "LOCAL_TEST_ONLY"),
        ({"STORAGE_PROVIDER": "s3"}, "LOCAL_TEST_ONLY"),
        ({"TENANT_B_CLERK_E2E_USER_ID": ""}, "MISSING_CLERK_IDENTITY"),
        ({"TENANT_B_CLERK_E2E_USER_ID": "user_a"}, "SAME_CLERK_USER"),
        ({"TENANT_B_CLERK_E2E_ORG_ID": "org_a"}, "SHARE_ONE_ORGANIZATION"),
    ],
)
def test_the_seed_fails_closed(override: dict[str, str], message: str) -> None:
    with pytest.raises(seed.SeedError, match=message):
        seed.require_environment(copy.deepcopy(GOOD_ENV) | override)


# ── runtime SHA receipt (verified against the process, not a URL) ────────────────────────────────────────────
def _fake_process(tmp_path: Path, command: str, cwd: Path) -> Path:
    proc = tmp_path / "proc"
    (proc / "4242").mkdir(parents=True)
    (proc / "4242" / "cmdline").write_bytes(command.replace(" ", "\0").encode())
    (proc / "4242" / "cwd").symlink_to(cwd)
    (proc / "self").mkdir()  # a non-numeric entry is ignored
    return proc


def _git_workspace(tmp_path: Path) -> tuple[Path, str]:
    import subprocess

    workspace = tmp_path / "ws"
    workspace.mkdir()
    env = {
        "GIT_AUTHOR_NAME": "t",
        "GIT_AUTHOR_EMAIL": "t@t",
        "GIT_COMMITTER_NAME": "t",
        "GIT_COMMITTER_EMAIL": "t@t",
        "PATH": os.environ["PATH"],
    }
    subprocess.run(["git", "init", "-q"], cwd=workspace, check=True, env=env)
    (workspace / "f").write_text("x")
    subprocess.run(["git", "add", "f"], cwd=workspace, check=True, env=env)
    subprocess.run(["git", "commit", "-q", "-m", "c"], cwd=workspace, check=True, env=env)
    sha = subprocess.run(
        ["git", "rev-parse", "HEAD"], cwd=workspace, check=True, capture_output=True, text=True
    ).stdout.strip()
    return workspace, sha


def test_the_receipt_matches_when_the_process_runs_the_candidate_sha(tmp_path: Path) -> None:
    from scripts import write_integral_runtime_receipt as receipt

    workspace, sha = _git_workspace(tmp_path)
    proc = _fake_process(tmp_path, "python -m uvicorn src.main:app --port 8000", workspace)
    result = receipt.build_receipt(
        role="backend",
        pattern="uvicorn src.main:app",
        expected_sha=sha,
        workspace=workspace,
        proc_root=proc,
    )
    assert (
        result["matches"] is True
        and result["runtime_sha"] == sha
        and result["cwd_inside_workspace"] is True
    )
    assert "cmdline" not in json.dumps(result)
    assert "uvicorn" not in json.dumps(result), "the command line is never recorded"


def test_the_receipt_fails_for_another_sha(tmp_path: Path) -> None:
    from scripts import write_integral_runtime_receipt as receipt

    workspace, _ = _git_workspace(tmp_path)
    proc = _fake_process(tmp_path, "uvicorn src.main:app", workspace)
    result = receipt.build_receipt(
        role="backend",
        pattern="uvicorn",
        expected_sha="0" * 40,
        workspace=workspace,
        proc_root=proc,
    )
    assert result["matches"] is False
    assert result["process_found"] is True


def test_the_receipt_fails_for_a_process_outside_the_workspace(tmp_path: Path) -> None:
    from scripts import write_integral_runtime_receipt as receipt

    workspace, sha = _git_workspace(tmp_path)
    elsewhere = tmp_path / "elsewhere"
    elsewhere.mkdir()
    proc = _fake_process(tmp_path, "uvicorn src.main:app", elsewhere)
    result = receipt.build_receipt(
        role="backend", pattern="uvicorn", expected_sha=sha, workspace=workspace, proc_root=proc
    )
    assert result["cwd_inside_workspace"] is False
    assert result["matches"] is False


def test_the_receipt_fails_when_there_is_no_such_process(tmp_path: Path) -> None:
    from scripts import write_integral_runtime_receipt as receipt

    workspace, sha = _git_workspace(tmp_path)
    proc = _fake_process(tmp_path, "something else", workspace)
    result = receipt.build_receipt(
        role="frontend", pattern="next dev", expected_sha=sha, workspace=workspace, proc_root=proc
    )
    assert result["process_found"] is False
    assert result["matches"] is False


def test_a_sha_from_the_environment_is_kept_only_when_it_is_one(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("GITHUB_SHA", "a" * 40)
    monkeypatch.setenv("FRONTEND_RUNTIME_SHA", "not a sha; rm -rf /")
    monkeypatch.delenv("BACKEND_RUNTIME_SHA", raising=False)
    assert verifier._sha_from_env("GITHUB_SHA") == "a" * 40
    assert verifier._sha_from_env("FRONTEND_RUNTIME_SHA") is None
    assert verifier._sha_from_env("BACKEND_RUNTIME_SHA") is None


def test_the_verifier_refuses_to_run_without_a_database_url(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.delenv("DATABASE_URL", raising=False)
    assert verifier.main() == 2
