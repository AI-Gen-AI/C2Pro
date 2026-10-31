"""Contract completion deadline vs schedule planned finish, end to end (DET-CRS-SCHDEAD).

TS-INT-COH-SCHDEAD-001. REAL PostgreSQL, a REAL ``.docx`` contract and a REAL ``.xlsx``
schedule uploaded through the real upload / reupload use cases, the REAL ingestion worker
``_process`` (real parsers, real clause persistence, real #711 processing authority), the REAL
schedule-finish loader and the REAL ``POST /coherence/evaluate`` endpoint function. Only object
storage (local), embeddings and the broker are replaced; no AI model is involved anywhere.

A finding needs two sourced, comparable, calendar-exact dates; every other case is
NOT_ASSESSED and raises no alert.
"""

from __future__ import annotations

import hashlib
import json
from contextlib import asynccontextmanager
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

import docx
import openpyxl
import pytest
from fastapi import UploadFile
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from src.coherence.router import CoherenceEvaluateRequest, evaluate_project_coherence
from src.coherence.schedule_clause_builder import load_schedule_planned_finish
from src.core.auth.models import SubscriptionPlan, Tenant
from src.core.tasks import ingestion_tasks
from src.documents.adapters.persistence.sqlalchemy_document_repository import (
    SqlAlchemyDocumentRepository,
)
from src.documents.adapters.rag import rag_service as rag_service_module
from src.documents.adapters.storage.local_file_storage_service import LocalFileStorageService
from src.documents.application.reupload_document_use_case import ReuploadDocumentUseCase
from src.documents.application.upload_document_use_case import UploadDocumentUseCase
from src.documents.domain.models import Document, DocumentType
from src.temporal.adapters.persistence.document_revision_repository import (
    SqlAlchemyDocumentRevisionRepository,
)
from src.temporal.adapters.persistence.project_event_repository import (
    SqlAlchemyProjectEventRepository,
)
from src.temporal.application import project_snapshot_trigger

pytestmark = pytest.mark.asyncio

RULE = "DET-CRS-SCHDEAD"
PROVISIONAL_RULE = "ADV-SCHDEAD-PROVISIONAL"
TRIPLET_DIR = (
    Path(__file__).resolve().parents[4] / "web" / "src" / "tests" / "e2e" / "test-data" / "triplet"
)
COMPLETION_CLAUSE = "La fecha de finalización de las obras será el 30 de junio de 2027."
DURATION_CLAUSE = "El plazo de ejecución de las obras será de dieciocho meses desde la firma."
PRICE_CLAUSE = "El precio del contrato asciende a 1.000.000 euros."

LATE_SCHEDULE = [
    ["Foundations", datetime(2027, 1, 4), datetime(2027, 3, 31)],
    ["Structure", datetime(2027, 4, 1), datetime(2027, 7, 30)],
    ["Handover", datetime(2027, 8, 2), datetime(2027, 9, 30)],
]
ON_TIME_SCHEDULE = [
    ["Foundations", datetime(2027, 1, 4), datetime(2027, 3, 31)],
    ["Handover", datetime(2027, 4, 1), datetime(2027, 5, 31)],
]
AMBIGUOUS_SCHEDULE = [["Handover", "01/04/2027", "30/09/2027"]]


def _docx(paragraphs: list[str]) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx(rows: list[list[Any]]) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Task", "Start Date", "End Date"])
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _dsn(db: AsyncSession) -> str:
    url = db.get_bind().url
    if url.get_driver_name() != "asyncpg":
        url = url.set(drivername="postgresql+asyncpg")
    return url.render_as_string(hide_password=False)


class _ProjectRepo:
    async def exists_by_id(self, _project_id: UUID, _tenant_id: UUID) -> bool:
        return True


async def _fake_embed(texts: list[str]) -> list[list[float]]:
    return [[0.001 * (i + 1)] * 1536 for i, _ in enumerate(texts)]


@pytest.fixture
async def worker(db: AsyncSession, monkeypatch: pytest.MonkeyPatch, tmp_path: Path):
    engine = create_async_engine(_dsn(db), pool_size=5, max_overflow=5)
    maker = async_sessionmaker(engine, expire_on_commit=False)

    @asynccontextmanager
    async def raw_session():
        async with maker() as session:
            try:
                yield session
            except Exception:
                await session.rollback()
                raise

    @asynccontextmanager
    async def tenant_session(tenant_id):
        async with maker() as session:
            await session.execute(
                text("SELECT set_config('app.current_tenant', :t, true)"), {"t": str(tenant_id)}
            )
            yield session
            await session.commit()

    async def _no_init() -> None:
        return None

    async def _no_heartbeat(**_: Any) -> None:
        return None

    class _Trigger:
        def __init__(self, **_: Any) -> None:
            pass

        async def execute(self, **_: Any) -> dict[str, Any]:
            return {"task_id": None, "task_name": "fake", "queue": "document_parsing"}

    storage = LocalFileStorageService(base_dir=tmp_path)
    monkeypatch.setattr(ingestion_tasks, "get_raw_session", raw_session)
    monkeypatch.setattr(ingestion_tasks, "init_db", _no_init)
    monkeypatch.setattr(ingestion_tasks, "build_storage_service", lambda: storage)
    monkeypatch.setattr(ingestion_tasks, "TriggerDocumentAnalysisUseCase", _Trigger)
    monkeypatch.setattr(ingestion_tasks, "_document_processing_heartbeat_loop", _no_heartbeat)
    monkeypatch.setattr(rag_service_module, "_embed_texts", _fake_embed)
    monkeypatch.setattr(project_snapshot_trigger, "get_session_with_tenant", tenant_session)
    monkeypatch.setattr(project_snapshot_trigger, "enqueue_project_snapshot", lambda **_: None)
    for module in ("upload_document_use_case", "reupload_document_use_case"):
        monkeypatch.setattr(
            f"src.documents.application.{module}.enqueue_project_snapshot", lambda **_: None
        )
    monkeypatch.setattr(
        "src.documents.application.reupload_document_use_case._enqueue_document_processing",
        lambda *_args, **_kwargs: None,
    )
    try:
        yield storage
    finally:
        await engine.dispose()


async def _project(db: AsyncSession, tenant_id: UUID) -> UUID:
    if await db.get(Tenant, tenant_id) is None:
        db.add(Tenant(id=tenant_id, name="t", slug=f"t-{tenant_id.hex[:8]}",
                      subscription_plan=SubscriptionPlan.PROFESSIONAL, ai_budget_monthly=100.0))
        await db.commit()
    project_id = uuid4()
    await db.execute(
        text("INSERT INTO projects (id, tenant_id, name, code, project_type, status, currency, "
             "created_at, updated_at) VALUES (:id, :tid, 'schdead', :code, 'construction', "
             "'active', 'EUR', now(), now())"),
        {"id": project_id, "tid": tenant_id, "code": f"P-{project_id.hex[:8]}"},
    )
    await db.commit()
    return project_id


async def _upload_and_ingest(
    db: AsyncSession, storage: Any, user: Any, project_id: UUID, *, filename: str,
    content: bytes, document_type: DocumentType,
) -> tuple[Document, Any]:
    revisions = SqlAlchemyDocumentRevisionRepository(db)
    document = await UploadDocumentUseCase(
        document_repository=SqlAlchemyDocumentRepository(db),
        storage_service=storage,
        project_repository=_ProjectRepo(),  # type: ignore[arg-type]
        revision_repository=revisions,
        event_repository=SqlAlchemyProjectEventRepository(db),
    ).execute(
        project_id=project_id,
        file=UploadFile(filename=filename, file=BytesIO(content)),
        document_type=document_type,
        user_id=user.id,
        tenant_id=user.tenant_id,
    )
    [revision] = await revisions.list_lineage(document.id, user.tenant_id)
    result = await ingestion_tasks._process(document.id, revision.revision_id)
    assert result["status"] == "success", result
    return document, revision


async def _metadata(db: AsyncSession, document_id: UUID) -> dict[str, Any]:
    await db.commit()
    row = (
        await db.execute(
            text("SELECT document_metadata FROM documents WHERE id = :d"), {"d": document_id}
        )
    ).one()
    return dict(row.document_metadata or {})


async def approve_document(
    db: AsyncSession, user: Any, project_id: UUID, document: Document, revision: Any
) -> None:
    """The human approval: a TRUSTED artifact bound to exactly ``revision`` (what HITL approval leaves)."""
    from src.analysis.adapters.persistence.document_artifact_repository import (
        SqlAlchemyDocumentArtifactRepository,
    )
    from src.analysis.domain.contracts import DocumentArtifact, RiskItem

    await SqlAlchemyDocumentArtifactRepository(db).save(
        DocumentArtifact(
            document_id=str(document.id),
            document_revision_id=str(revision.revision_id),
            doc_type="contract",
            extracted_risks=[RiskItem(title="r", description="d")],
        ),
        project_id=project_id,
        tenant_id=user.tenant_id,
    )
    await db.commit()


_trust = approve_document  # name used by the main-line revision-authority tests


async def _scenario(
    db: AsyncSession, worker: Any, user: Any, *, contract: list[str], schedule: list[list[Any]],
    approve_contract: bool = False,
) -> tuple[UUID, Document, Any, Document, Any]:
    project_id = await _project(db, user.tenant_id)
    contract_doc, contract_rev = await _upload_and_ingest(
        db, worker, user, project_id, filename="contract.docx",
        content=_docx(contract), document_type=DocumentType.CONTRACT,
    )
    schedule_doc, schedule_rev = await _upload_and_ingest(
        db, worker, user, project_id, filename="schedule.xlsx",
        content=_xlsx(schedule), document_type=DocumentType.SCHEDULE,
    )
    if approve_contract:
        await approve_document(db, user, project_id, contract_doc, contract_rev)
    return project_id, contract_doc, contract_rev, schedule_doc, schedule_rev


async def _evaluate(db: AsyncSession, user: Any, project_id: UUID) -> Any:
    return await evaluate_project_coherence(
        payload=CoherenceEvaluateRequest(project_id=project_id),
        include_diagnostics=True,
        db=db,
        current_user=user,
        flags_service=None,
    )


def _schdead(result: Any) -> list[Any]:
    """The provisional advisory (the only SCHDEAD-related alert a submitted schedule can raise)."""
    assert not [a for a in result.alerts if a.rule_id == RULE], "never a confirmed SCHDEAD alert"
    assert not [s for s in result.finding_signals if s.rule_id == RULE], "never a scoring signal"
    return [alert for alert in result.alerts if alert.rule_id == PROVISIONAL_RULE]


# ── the worker records the observation, with identity ───────────────────────


async def test_worker_records_the_schedule_planned_finish_with_its_row_identity(
    db: AsyncSession, worker: Any, test_user: Any
) -> None:
    _p, _c, _cr, schedule_doc, schedule_rev = await _scenario(
        db, worker, test_user, contract=[COMPLETION_CLAUSE], schedule=LATE_SCHEDULE
    )

    finish = (await _metadata(db, schedule_doc.id))["schedule_planned_finish"]

    assert finish["date"] == "2027-09-30"
    assert finish["source_locator"] == "Sheet!4"
    assert finish["activity"] == "Handover"
    assert finish["document_id"] == str(schedule_doc.id)
    assert finish["revision_id"] == str(schedule_rev.revision_id)
    assert finish["basis"] == "SOURCE_OBSERVED_NOT_APPROVED_BASELINE"


async def test_ambiguous_schedule_dates_record_a_reason_and_no_date(
    db: AsyncSession, worker: Any, test_user: Any
) -> None:
    _p, _c, _cr, schedule_doc, _sr = await _scenario(
        db, worker, test_user, contract=[COMPLETION_CLAUSE], schedule=AMBIGUOUS_SCHEDULE
    )

    finish = (await _metadata(db, schedule_doc.id))["schedule_planned_finish"]

    assert "date" not in finish
    assert finish["reason"] == "no_row_with_unambiguous_start_and_end"


# ── evaluate: finding, no finding, NOT_ASSESSED ─────────────────────────────


async def test_late_schedule_raises_a_provisional_advisory_pointing_at_the_contract_clause(
    db: AsyncSession, worker: Any, test_user: Any
) -> None:
    project_id, contract_doc, contract_rev, schedule_doc, schedule_rev = await _scenario(
        db, worker, test_user, contract=[PRICE_CLAUSE, COMPLETION_CLAUSE], schedule=LATE_SCHEDULE
    )

    result = await _evaluate(db, test_user, project_id)

    [alert] = _schdead(result)
    assert str(alert.category).lower() == "schedule"
    assert str(alert.severity).lower() == "medium"  # a risk to look at, not a confirmed breach
    assert "not a confirmed breach" in alert.message
    assert "not an approved temporal baseline" in alert.message
    assert "2027-09-30" in alert.message and "2027-06-30" in alert.message
    # Both sources are named: the contract clause (link) and the schedule document/revision/row.
    claim = alert.evidence.claim
    assert str(schedule_doc.id) in claim
    assert str(schedule_rev.revision_id) in claim
    assert "Sheet!4" in claim
    clause_text = (
        await db.execute(
            text("SELECT full_text, document_id::text AS document_id FROM clauses WHERE id = :c"),
            {"c": UUID(alert.evidence.source_clause_id)},
        )
    ).one()
    assert "30 de junio de 2027" in clause_text.full_text
    assert clause_text.document_id == str(contract_doc.id)


async def test_the_provisional_advisory_does_not_move_the_score_or_the_scoring_signals(
    db: AsyncSession, worker: Any, test_user: Any
) -> None:
    """Same contract, same documents: the late schedule adds an advisory but not one scoring signal."""
    late_project, *_ = await _scenario(
        db, worker, test_user, contract=[PRICE_CLAUSE, COMPLETION_CLAUSE], schedule=LATE_SCHEDULE
    )
    ontime_project, *_ = await _scenario(
        db, worker, test_user, contract=[PRICE_CLAUSE, COMPLETION_CLAUSE], schedule=ON_TIME_SCHEDULE
    )

    late = await _evaluate(db, test_user, late_project)
    ontime = await _evaluate(db, test_user, ontime_project)

    assert _schdead(late) != [] and _schdead(ontime) == []
    assert late.overall_score == ontime.overall_score
    assert sorted(s.rule_id for s in late.finding_signals) == sorted(
        s.rule_id for s in ontime.finding_signals
    )


async def test_the_provisional_advisory_is_persisted_with_its_navigable_source(
    db: AsyncSession, worker: Any, test_user: Any
) -> None:
    """A scoreless /evaluate itself must commit the provisional advisory to Alerts."""
    project_id, contract_doc, *_ = await _scenario(
        db, worker, test_user, contract=[PRICE_CLAUSE, COMPLETION_CLAUSE], schedule=LATE_SCHEDULE,
        approve_contract=True,  # the document link is only authorised by an approved contract (#1111)
    )
    result = await _evaluate(db, test_user, project_id)
    assert result.overall_score is None, "this regression must exercise the score-withheld path"
    [advisory] = _schdead(result)

    # Do NOT call the internal mirror directly: persistence is part of the API flow.
    # Read after /evaluate's transaction boundary, not from an uncommitted fixture.
    scored_rows = (
        await db.execute(
            text("SELECT count(*) FROM coherence_results WHERE project_id = :p"),
            {"p": project_id},
        )
    ).scalar_one()
    assert scored_rows == 0, "scoreless findings must not fabricate a scored snapshot"

    row = (
        await db.execute(
            text("SELECT category::text AS category, severity::text AS severity, "
                 "source_clause_id::text AS clause, alert_metadata, affected_entities "
                 "FROM alerts WHERE project_id = :p AND rule_id = :r"),
            {"p": project_id, "r": PROVISIONAL_RULE},
        )
    ).one()
    assert row.category.upper() == "TIME"  # canonical persisted taxonomy for schedule/time alerts
    assert row.severity.lower() == "medium"
    assert row.clause == advisory.evidence.source_clause_id  # navigation to the contract clause
    assert row.alert_metadata["detection_evidence"]["source_document_id"] == str(contract_doc.id)
    assert "Sheet!4" in row.alert_metadata["detection_evidence"]["claim"]  # the XLSX row is named in the evidence
    assert str(contract_doc.id) in row.affected_entities["documents"]

    # A repeated scoreless evaluation must update the same durable finding, not
    # create a second alert or fabricate a numeric score.
    repeated = await _evaluate(db, test_user, project_id)
    assert repeated.overall_score is None
    alert_rows = (
        await db.execute(
            text("SELECT count(*) FROM alerts WHERE project_id = :p AND rule_id = :r"),
            {"p": project_id, "r": PROVISIONAL_RULE},
        )
    ).scalar_one()
    assert alert_rows == 1


async def test_a_schedule_of_another_project_never_feeds_this_projects_comparison(
    db: AsyncSession, worker: Any, test_user: Any
) -> None:
    project_id = await _project(db, test_user.tenant_id)
    await _upload_and_ingest(
        db, worker, test_user, project_id, filename="contract.docx",
        content=_docx([PRICE_CLAUSE, COMPLETION_CLAUSE]), document_type=DocumentType.CONTRACT,
    )
    other_project = await _project(db, test_user.tenant_id)
    await _upload_and_ingest(
        db, worker, test_user, other_project, filename="schedule.xlsx",
        content=_xlsx(LATE_SCHEDULE), document_type=DocumentType.SCHEDULE,
    )

    assert await load_schedule_planned_finish(db, project_id, test_user.tenant_id) is None
    assert _schdead(await _evaluate(db, test_user, project_id)) == []


async def test_on_time_schedule_raises_no_alert(
    db: AsyncSession, worker: Any, test_user: Any
) -> None:
    project_id, *_ = await _scenario(
        db, worker, test_user, contract=[COMPLETION_CLAUSE], schedule=ON_TIME_SCHEDULE
    )

    assert _schdead(await _evaluate(db, test_user, project_id)) == []


async def test_contract_with_only_a_duration_is_not_assessed_not_flagged(
    db: AsyncSession, worker: Any, test_user: Any
) -> None:
    project_id, *_ = await _scenario(
        db, worker, test_user, contract=[PRICE_CLAUSE, DURATION_CLAUSE], schedule=LATE_SCHEDULE
    )

    assert _schdead(await _evaluate(db, test_user, project_id)) == []


async def test_a_contract_stating_two_completion_dates_is_not_assessed_not_flagged(
    db: AsyncSession, worker: Any, test_user: Any
) -> None:
    """The first date must never win: a revised completion date in the same clause is a conflict."""
    conflicting = (
        "The completion date is 2027-06-30. The contractual completion date is revised to "
        "2027-12-31."
    )
    project_id, *_ = await _scenario(
        db, worker, test_user, contract=[PRICE_CLAUSE, conflicting], schedule=LATE_SCHEDULE
    )

    assert _schdead(await _evaluate(db, test_user, project_id)) == []


async def test_schedule_with_ambiguous_dates_is_not_assessed_not_flagged(
    db: AsyncSession, worker: Any, test_user: Any
) -> None:
    project_id, *_ = await _scenario(
        db, worker, test_user, contract=[COMPLETION_CLAUSE], schedule=AMBIGUOUS_SCHEDULE
    )

    assert _schdead(await _evaluate(db, test_user, project_id)) == []


async def test_a_schedule_revision_that_is_not_current_is_not_assessed(
    db: AsyncSession, worker: Any, test_user: Any
) -> None:
    """V2 uploaded but not yet ingested/trusted: V1's observation is no longer current."""
    project_id, _c, _cr, schedule_doc, _sr = await _scenario(
        db, worker, test_user, contract=[COMPLETION_CLAUSE], schedule=LATE_SCHEDULE
    )
    assert _schdead(await _evaluate(db, test_user, project_id)) != []

    revisions = SqlAlchemyDocumentRevisionRepository(db)
    await ReuploadDocumentUseCase(
        document_repository=SqlAlchemyDocumentRepository(db),
        revision_repository=revisions,
        storage_service=worker,
        event_repository=SqlAlchemyProjectEventRepository(db),
    ).execute(
        tenant_id=test_user.tenant_id,
        document_id=schedule_doc.id,
        file_content=_xlsx(ON_TIME_SCHEDULE),
        user_id=test_user.id,
    )

    loaded = await load_schedule_planned_finish(db, project_id, test_user.tenant_id)
    assert loaded is not None
    assert loaded["reason"] == "schedule_revision_not_current"
    assert "date" not in loaded
    assert _schdead(await _evaluate(db, test_user, project_id)) == []


# ── the browser journey's fixture files, through the same backend path ──────


async def test_the_triplet_fixture_files_raise_the_budget_alerts_and_a_provisional_schedule_risk(
    db: AsyncSession, worker: Any, test_user: Any
) -> None:
    """Grounds the E2E journey: its real fixture files, real worker, real evaluate; no AI."""
    manifest = json.loads((TRIPLET_DIR / "triplet.manifest.json").read_text(encoding="utf-8"))
    for name, expected in manifest["sha256"].items():
        assert hashlib.sha256((TRIPLET_DIR / name).read_bytes()).hexdigest() == expected, name

    project_id = await _project(db, test_user.tenant_id)
    schedule_doc = schedule_rev = None
    for filename, document_type in (
        ("triplet-contract.docx", DocumentType.CONTRACT),
        ("triplet-budget.xlsx", DocumentType.BUDGET),
        ("triplet-schedule.xlsx", DocumentType.SCHEDULE),
        ("triplet-specification.docx", DocumentType.TECHNICAL_SPEC),
    ):
        document, revision = await _upload_and_ingest(
            db, worker, test_user, project_id, filename=filename,
            content=(TRIPLET_DIR / filename).read_bytes(), document_type=document_type,
        )
        if document_type is DocumentType.SCHEDULE:
            schedule_doc, schedule_rev = document, revision

    finish = (await _metadata(db, schedule_doc.id))["schedule_planned_finish"]
    assert finish["date"] == manifest["schedule"]["planned_finish"]
    assert finish["source_locator"] == manifest["schedule"]["source_locator"]
    assert finish["revision_id"] == str(schedule_rev.revision_id)

    result = await _evaluate(db, test_user, project_id)

    rules = {alert.rule_id for alert in result.alerts}
    assert rules >= set(manifest["expected_rules"])
    assert rules >= set(manifest["expected_provisional"])
    assert RULE not in rules  # a submitted schedule never confirms a breach
    [advisory] = _schdead(result)
    assert manifest["contract"]["completion_deadline"] in advisory.message

    # Four families, still no valid coverage: the score is withheld AND the findings are durable.
    assert result.overall_score is None
    await db.commit()
    persisted = {
        row.rule_id
        for row in (
            await db.execute(
                text("SELECT rule_id FROM alerts WHERE project_id = :p"), {"p": project_id}
            )
        ).all()
    }
    assert persisted >= set(manifest["expected_rules"]) | set(manifest["expected_provisional"])
