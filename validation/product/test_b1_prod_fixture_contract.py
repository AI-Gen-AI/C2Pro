"""Contracts for bounded #867 B1 production fixtures."""

from __future__ import annotations

import importlib.util
from pathlib import Path

import openpyxl

from src.coherence.cross_document.assembly import cross_document_signals
from src.coherence.models import Clause
from src.coherence.rules_engine.deterministic import RetentionRateEvaluator

REPO_ROOT = Path(__file__).resolve().parents[2]
GENERATOR = REPO_ROOT / "apps/api/scripts/generate_b1_prod_triplet.py"


def _module():
    spec = importlib.util.spec_from_file_location("b1_fixture_generator", GENERATOR)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_b1_fixture_has_distinct_genuine_and_false_positive_truth() -> None:
    module = _module()
    assert module.NET_BUDGET_ROWS_TOTAL == module.CONTRACT_NET_EUR
    assert module.BUDGET_GROSS_EUR == module.CONTRACT_NET_EUR * (1 + module.VAT_RATE)
    assert "excluding VAT" in module.CONTRACT_SOURCE
    assert "fifteen percent (15%)" in module.CONTRACT_SOURCE
    assert module.SCHEDULE_ROWS[-1][2] == "2025-12-31"


def test_b1_genuine_retention_finding_is_real_deterministic_output() -> None:
    finding = RetentionRateEvaluator().evaluate_v3(
        Clause(
            id="synthetic-retention",
            text="Retention shall be fifteen percent (15%).",
            data={"retention_pct": 15.0},
        )
    )
    assert finding is not None
    assert finding.rule_id == "DET-BUD-RETENTION"
    assert finding.category == "BUDGET"


def test_b1_vat_basis_produces_expected_cross_document_false_positive_candidates() -> None:
    module = _module()
    signals = cross_document_signals(
        [
            Clause(
                id="contract-total",
                text="Contract price excludes VAT",
                data={"contract_total": module.CONTRACT_NET_EUR, "currency": "EUR"},
            ),
            Clause(
                id="budget-total",
                text="Budget total includes VAT",
                data={"stated_total": module.BUDGET_GROSS_EUR, "currency": "EUR"},
            ),
        ]
    )
    rule_ids = {signal.rule_id for signal in signals}
    assert "DET-CRS-CONBUD" in rule_ids
    assert "DET-CRS-NEGMARGIN" in rule_ids


def test_b1_budget_xlsx_uses_canonical_parser_headers(tmp_path: Path) -> None:
    module = _module()
    path = tmp_path / "budget.xlsx"
    module.generate_budget(path)
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        assert headers == ["Item", "Quantity", "Unit Price", "Total", "Unit"]
        assert sheet.max_row >= 3
        declared_total = sheet.cell(row=sheet.max_row, column=4).value
        assert declared_total == module.BUDGET_GROSS_EUR
    finally:
        workbook.close()


def test_b1_schedule_xlsx_uses_canonical_parser_headers(tmp_path: Path) -> None:
    module = _module()
    path = tmp_path / "schedule.xlsx"
    module.generate_schedule(path)
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        assert headers == [
            "Task",
            "Start Date",
            "End Date",
            "Duration",
            "WBS",
            "Predecessors",
        ]
        assert sheet.max_row >= 3
        assert sheet.cell(row=sheet.max_row, column=3).value == "2025-12-31"
    finally:
        workbook.close()


def test_b1_contract_pdf_is_text_layer_and_declares_tax_basis(tmp_path: Path) -> None:
    module = _module()
    path = tmp_path / "contract.pdf"
    module.generate_contract(path)
    import fitz

    with fitz.open(path) as document:
        text = "\n".join(page.get_text() for page in document)
    assert "EUR 2,400,000 excluding VAT" in text
    assert "fifteen percent (15%)" in text
    assert "synthetic" in text.lower()
